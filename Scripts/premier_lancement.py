#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""premier_lancement.py — cree TON profil a partir de ton CV.

    python Scripts/premier_lancement.py chemin/vers/mon_cv.pdf      (ou .docx)
    python Scripts/premier_lancement.py mon_cv.docx --auto           (aucune question)

Ce que fait le programme :
  1. lit le CV (PDF ou Word) : nom, email, telephone, LinkedIn, ville, langues,
     competences reconnues (catalogue_competences.py) ;
  2. te pose cinq questions : metiers vises, annees d'experience, villes,
     accords (il/elle), listes d'entreprises ;
  3. ecrit profil/profil.json — tout ce que le moteur sait de toi —, copie le CV
     dans CV/Modeles/, prepare sources_config.json (mots-cles de recherche).

Un CV .docx est PREFERABLE : les CV generes par offre partent de ce modele et
peuvent y inserer l'intitule du poste (ecris « [nom du poste] » dans ton
accroche). Un CV .pdf est envoye tel quel.

Rien n'est envoye nulle part : tout reste dans ce dossier.
Ensuite : relis profil/profil.json (les champs « A COMPLETER » surtout).
"""
import argparse
import json
import re
import shutil
import sys
import unicodedata
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = Path(__file__).resolve().parent
RACINE = HERE.parent
sys.path.insert(0, str(HERE))

from catalogue_competences import FAMILLES, catalogue   # noqa: E402

A_COMPLETER = "A COMPLETER"

# ── 1. Lecture du CV ─────────────────────────────────────────────────────────


def lire_cv(chemin):
    p = Path(chemin)
    if p.suffix.lower() == ".pdf":
        import fitz
        with fitz.open(str(p)) as d:
            return "\n".join(page.get_text() for page in d)
    if p.suffix.lower() == ".docx":
        from docx import Document
        d = Document(str(p))
        lignes = [x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    lignes.append(c.text)
        return "\n".join(lignes)
    raise SystemExit(f"Format non pris en charge : {p.suffix} (PDF ou DOCX)")


_RX_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_RX_TEL = re.compile(r"(?:\+\d{2,3}[\s.]?\(?0?\)?[\s.]?[1-9](?:[\s.-]?\d{2}){4}|\b0[1-9](?:[\s.-]?\d{2}){4})\b")
_RX_LINKEDIN = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[\w%-]+/?", re.I)
_RX_TITRE_SECTION = re.compile(
    r"^\s*(profil|profile|r[ée]sum[ée]|summary|about( me)?|[àa] propos|objectif|pr[ée]sentation)\s*:?\s*$",
    re.I)
_RX_DIPLOME = re.compile(r"\b(master|licence|bachelor|bts|dut|but|mba|doctorat|phd|ing[ée]nieur|bac\s*\+\s*\d)",
                         re.I)


def _sans_accent(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if not unicodedata.combining(c))


def identite_du_cv(texte):
    lignes = [l.strip() for l in texte.splitlines() if l.strip()]
    email = (_RX_EMAIL.search(texte) or [""])[0]
    tel = (_RX_TEL.search(texte) or [""])[0]
    lk = (_RX_LINKEDIN.search(texte) or [""])[0]
    prenom = nom = ""
    # Le nom : une des premieres lignes, 2 a 4 mots alphabetiques, sans chiffre.
    for l in lignes[:8]:
        mots = l.replace("-", " - ").split()
        if (2 <= len([m for m in mots if m != "-"]) <= 4 and not re.search(r"[\d@/:|·•]", l)
                and all(re.fullmatch(r"[A-Za-zÀ-ÿ'’-]+", m) for m in mots)):
            vrais = [m for m in l.split() if m]
            # Le mot en MAJUSCULES est le nom de famille, s'il y en a un.
            maj = [m for m in vrais if m.isupper() and len(m) > 1]
            if maj and len(maj) < len(vrais):
                nom = " ".join(m.capitalize() for m in maj)
                prenom = " ".join(m for m in vrais if m not in maj)
            else:
                prenom, nom = vrais[0], " ".join(vrais[1:])
            break
    # La ville : un secteur du catalogue des zones nomme dans l'en-tete.
    ville = ""
    entete = _sans_accent(" ".join(lignes[:12])).lower()
    import zones
    for cle, s in zones.CATALOGUE.items():
        for lieu in s.get("lieux", ()):
            if len(lieu) > 3 and re.search(r"(?<![a-z])" + re.escape(lieu) + r"(?![a-z])", entete):
                ville = lieu.title()
                break
        if ville:
            break
    return {"prenom": prenom, "nom": nom, "email": email, "telephone": tel,
            "linkedin": lk, "ville": ville}


def resume_du_cv(texte):
    """Le paragraphe sous « Profil / Résumé / Summary », s'il existe."""
    lignes = texte.splitlines()
    for i, l in enumerate(lignes):
        if _RX_TITRE_SECTION.match(l):
            bloc = []
            for x in lignes[i + 1:i + 12]:
                x = x.strip()
                if not x:
                    if bloc:
                        break
                    continue
                if x.isupper() and len(x) < 40:          # titre de la section suivante
                    break
                bloc.append(x)
            txt = re.sub(r"\s+", " ", " ".join(bloc)).strip()
            if len(txt) > 60:
                return txt[:700]
    return ""


def competences_du_cv(texte):
    t = " " + texte.lower() + " "
    trouvees = {}
    for lab, v in catalogue().items():
        for s in v["synonymes"]:
            s = s.lower()
            ok = (re.search(r"\b" + re.escape(s.strip()) + r"\b", t) if len(s.strip()) <= 4 else s in t)
            if ok:
                trouvees[lab] = v
                break
    return trouvees


def langue_du_cv(texte):
    t = texte.lower()
    fr = len(re.findall(r"\b(et|les|des|une|pour|avec|exp[ée]rience|comp[ée]tences)\b", t))
    en = len(re.findall(r"\b(and|the|with|for|experience|skills)\b", t))
    return "en" if en > fr * 1.3 else "fr"


# ── 2. Questions ─────────────────────────────────────────────────────────────

def demander(question, defaut, auto, impose=None):
    # Une reponse donnee en option (--metiers...) passe avant la question : un agent
    # comme Claude Code renseigne tout sans clavier.
    if impose is not None:
        print(f"  {question} -> {impose}")
        return impose
    # Sans clavier (agent, tuyau), input() leverait EOFError : on prend la valeur proposee.
    if auto or not sys.stdin.isatty():
        print(f"  {question} -> {defaut}")
        return defaut
    try:
        r = input(f"  {question} [{defaut}] : ").strip()
    except EOFError:            # console sans clavier : la valeur proposee
        r = ""
    return r or defaut


_VIDES = {"junior", "senior", "stage", "stagiaire", "alternance", "alternant", "cdi", "cdd", "h/f", "f/h",
          "h-f", "f-h", "(h/f)", "de", "des", "du", "en", "et", "la", "le", "les", "d", "l", "au", "aux",
          "pour", "a", "à", "or", "and", "of", "the", "in", "vie", "v.i.e", "freelance"}


def mots_significatifs(titre):
    mots = re.findall(r"[a-zA-ZÀ-ÿ&+#.]+", titre.lower())
    return [m for m in mots if m not in _VIDES and len(m) >= 3]


def slug(s):
    s = _sans_accent(s)
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_") or "Domaine"


# ── 3. Construction du profil ────────────────────────────────────────────────

# Roles ecartes par defaut (rarement voulus par erreur), avec leur famille :
# une famille VISEE n'est jamais ecartee.
_EXCLUS_DEFAUT = [
    ("brand", "marketing"), ("re:design(er)?\\b", "design"), ("re:\\bux\\b", "design"),
    ("re:\\bui\\b", "design"), ("graphi", "design"), ("marketing", "marketing"), ("growth", "marketing"),
    ("sales", "commercial"), ("account executive", "commercial"), ("account manager", "commercial"),
    ("business developer", "commercial"), ("re:commercial\\b", "commercial"),
    ("talent acquisition", "rh"), ("recruiter", "rh"), ("recruteur", "rh"), ("people ops", "rh"),
    ("office manager", ""), ("customer success", "commercial"), ("customer care", "commercial"),
    ("customer support", "commercial"), ("community manager", "marketing"), ("copywriter", "marketing"),
    ("content manager", "marketing"), ("content creator", "marketing"), ("social media", "marketing"),
    ("juriste", "juridique"), ("legal counsel", "juridique"), ("avocat", "juridique"),
    ("re:paie\\b", "rh"), ("payroll", "rh"),
]


def familles_visees(metiers, comps):
    """Score par famille : intitules vises (fort) + competences du CV."""
    score = {f: 0 for f in FAMILLES}
    txt = " " + " ".join(metiers).lower() + " "
    for f, (_lib, mots) in FAMILLES.items():
        score[f] += 3 * sum(1 for m in mots if m.lower() in txt)
    for v in comps.values():
        if v.get("famille"):
            score[v["famille"]] += 1
    return sorted((f for f in score if score[f] > 0), key=lambda f: -score[f]), score


def niveau(annees):
    if annees <= 2:
        return {"mots_junior": ["junior", "graduate", "débutant", "debutant", "jeune diplôm", "alternance",
                                "apprenti", "stage", "stagiaire", "entry level", "0 à 2 ans", "1 à 2 ans"],
                "mots_senior": ["senior", "director", "directeur", "manager", "head of", "principal",
                                "vp ", "chef de", "expert confirmé", "expert ", "confirmé", "expérimenté"],
                "senior_en_fin": ["re:\\b(lead|head|principal|manager|director)\\b(?!\\s+(generation|gen)\\b)"],
                "mots_negatifs": [], "malus_experience": {"3": 1, "5": 3, "7": 4, "10": 5},
                "experience_max": 5}, True
    if annees <= 6:
        return {"mots_junior": [], "bonus_junior": 0,
                "mots_senior": ["director", "directeur", "directrice", "head of", "chief", "vp "],
                "malus_senior": 1, "senior_en_fin": [], "mots_negatifs": [],
                "malus_experience": {"9": 2, "12": 4}, "experience_max": 12}, False
    return {"mots_junior": [], "bonus_junior": 0, "mots_senior": [], "malus_senior": 0,
            "senior_en_fin": [], "mots_negatifs": [], "malus_experience": {}, "experience_max": 99}, False


def construire(cv_txt, cv_nom, rep):
    ident = rep["identite"]
    comps = rep["competences"]
    metiers = rep["metiers"]
    resume = resume_du_cv(cv_txt)
    ordre, score = familles_visees(metiers, comps)
    coeur = ordre[0] if ordre else ""
    mots_cibles = sorted({m for t in metiers for m in mots_significatifs(t)})

    # Competences du profil = celles du CV. La lettre cite des competences
    # que l'annonce demande ET que le CV porte : jamais une qu'il n'a pas.
    competences = {}
    for lab, v in comps.items():
        e = {k: v[k] for k in ("synonymes", "en", "prose_fr", "prose_en") if k in v}
        if v.get("technique") is False or not v.get("famille"):
            e["technique"] = False
        competences[lab] = e

    axes = {}
    for f in ordre:
        labs = [l for l, v in comps.items() if v.get("famille") == f]
        if not labs:
            continue
        axes[f] = {"poids": 1.0, "competences": labs}
    if len(ordre) > 1 and ordre[1] in axes:
        axes[ordre[1]]["poids"] = 2.0
    if coeur in axes:
        axes[coeur]["poids"] = 3.0

    niv, junior = niveau(rep["annees"])
    visees = {f for f in ordre if score[f] >= 3} | ({coeur} if coeur else set())
    exclus = [m for m, fam in _EXCLUS_DEFAUT if fam not in visees and not any(
        (m[3:] if m.startswith("re:") else m) in t.lower() for t in metiers)]

    domaines = []
    for t in metiers:
        domaines.append({
            "cle": slug(t)[:40], "libelle": t[:30], "mots_titre": mots_significatifs(t) or [t.lower()],
            "cv": cv_nom, "cv_en": "",
            "parcours_fr": resume, "parcours_en": "",
            "accroche_fr": t.lower(), "accroche_en": t.lower(),
            "preuve_fr": "", "preuve_en": "", "preuve_detente_fr": "",
        })
    domaines.append({"cle": "General", "libelle": "Autre", "mots_titre": [], "cv": cv_nom, "cv_en": "",
                     "parcours_fr": resume, "parcours_en": "",
                     "accroche_fr": (metiers[0].lower() if metiers else "candidat"),
                     "accroche_en": (metiers[0].lower() if metiers else "candidate"),
                     "preuve_fr": "", "preuve_en": "", "preuve_detente_fr": ""})

    # Forme masculine aussi : les annonces ecrivent « Chargé(e) », « Chargé H/F ».
    formes = []
    for t in metiers:
        formes.append(t.lower())
        masc = re.sub(r"(\w)ée\b", r"\1é", t.lower())
        if masc != t.lower():
            formes.append(masc)
    mots_cles = list(dict.fromkeys(formes + ([f"{t} junior" for t in formes[:2]] if junior else [])))
    top = sorted((l for l, v in comps.items() if v.get("famille")),
                 key=lambda l: -(comps[l].get("famille") == coeur))[:3]

    return {
        "_lisez_moi": ("Ton profil. Genere par premier_lancement.py a partir de ton CV : RELIS-LE. "
                       "Les champs vides ou « A COMPLETER » rendent les lettres plus pauvres. "
                       "Modele complet et commente : profil/exemple_finance_esg.json."),
        "identite": {**ident, "genre": rep["genre"]},
        "domaines": domaines,
        "competences": competences,
        "competences_de_base": top,
        "termes": {},
        "scoring": {
            "axes": axes, "axe_coeur": coeur if coeur in axes else (next(iter(axes), "")),
            "axe_complement": "", "cible_score": 3.0, "saturation_axe": 3,
            "plafond_sans_coeur": 6.0, "plafond_complement": 8.0, "densite_coeur_min": 0.0,
            "densite_coeur_longueur": 800,
            "intitules_coeur": mots_cibles, "intitules_coeur_forts": [t.lower() for t in metiers],
            "intitules_second": [], "plafond_second": 8.0,
            "veto_sigles": [], "veto_si": [], "veto_et": [], "plafond_veto": 6.0,
            "intitules_techniques": [], "intitules_techniques_sauf": [], "intitules_hors_metier": [],
        },
        "niveau": niv,
        "filtre_titre": {
            "metier": mots_cibles + [m for m in (FAMILLES.get(coeur, ("", []))[1])],
            "exclus": exclus, "exclus_sauf_complement": [], "complement": [], "lexique_large": [],
        },
        "tri": {"retirer_seniors": junior, "retirer_etudiants": True,
                "metiers_durs": [], "sauvetage": mots_cibles,
                "outils_max": 0 if coeur in ("dev", "data") else 4},
        # Langues parlees : le francais (outil en francais), plus celles du CV.
        # Une annonce redigee ou exigee dans une autre langue est plafonnee.
        "langues_parlees": ["fr"] + [c for nom, c in (("Anglais", "en"), ("Allemand", "de"),
                                                     ("Espagnol", "es"), ("Italien", "it"))
                                     if nom in comps],
        "zones": {"prioritaires": rep["prioritaires"], "secondaires": rep["secondaires"]},
        "collecte": {"mots_cles": mots_cles},
        "message": {"diplome_fr": "", "diplome_detente_fr": "", "diplome_en": "",
                    "langues_fr": "", "langues_detente_fr": "", "langues_en": "",
                    "sources": {}, "faits": {}, "phrase_contexte_fr": {}, "phrase_contexte_en": {}},
        "lettre": {"valeurs_entreprises_fr": {}, "valeurs_entreprises_en": {},
                   "phrase_profil_fr": "", "phrase_profil_en": "", "contextes": [],
                   "presentation_llm": f"Profil : {', '.join(metiers)}. {resume[:300]}"},
        "formulaire": {"pays": "France", "langues": ", ".join(
            l for l, v in comps.items() if v.get("technique") is False and not v.get("famille"))},
    }


def familles_tri_a_exclure(metiers):
    """Familles du tri (tri_file._METIERS) qui ne visent AUCUN des metiers
    choisis : elles sont retirees de la file. Une famille visee reste."""
    import tri_file
    exclure = []
    for fam, rx in tri_file._METIERS_CATALOGUE.items():
        if not any(rx.search(t) for t in metiers):
            exclure.append(fam)
    return exclure


# ── 4. Programme ─────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description="Cree profil/profil.json a partir de ton CV.")
    ap.add_argument("cv", help="ton CV, .pdf ou .docx")
    ap.add_argument("--auto", action="store_true", help="accepter toutes les reponses proposees")
    ap.add_argument("--force", action="store_true", help="remplacer un profil existant")
    ap.add_argument("--profil", default=str(RACINE / "profil" / "profil.json"))
    ap.add_argument("--metiers", help="metiers vises, separes par des virgules")
    ap.add_argument("--annees", help="annees d'experience, stages compris")
    ap.add_argument("--genre", choices=["m", "f"], help="accords des lettres")
    ap.add_argument("--prioritaires", help="secteurs prioritaires (zones.CATALOGUE), virgules")
    ap.add_argument("--secondaires", help="secteurs secondaires, virgules (vide = aucun)")
    ap.add_argument("--entreprises", choices=["o", "n"],
                    help="garder (o) ou vider (n) les entreprises finance/ESG suivies")
    a = ap.parse_args()

    sortie = Path(a.profil)
    if sortie.exists() and not a.force:
        print(f"Un profil existe deja : {sortie}\nRelance avec --force pour le remplacer "
              "(l'ancien est garde en .bak).")
        return 1

    cv = Path(a.cv)
    if not cv.exists():
        print(f"CV introuvable : {cv}")
        return 1
    print(f"\n1. Lecture du CV : {cv.name}")
    txt = lire_cv(cv)
    if len(txt.strip()) < 200:
        print("  ! Presque aucun texte lu. Un CV scanne (image) ne se lit pas : fournis un PDF texte ou un .docx.")
    ident = identite_du_cv(txt)
    comps = competences_du_cv(txt)
    print(f"  nom : {ident['prenom']} {ident['nom']} · email : {ident['email'] or '?'} · "
          f"tel : {ident['telephone'] or '?'} · ville : {ident['ville'] or '?'}")
    print(f"  {len(comps)} competence(s) reconnue(s) : {', '.join(sorted(comps))[:300]}")

    print("\n2. Quelques questions (Entree = valeur proposee)")
    ident["prenom"] = demander("Prénom", ident["prenom"] or A_COMPLETER, a.auto)
    ident["nom"] = demander("Nom", ident["nom"] or A_COMPLETER, a.auto)
    ident["email"] = demander("Email", ident["email"] or A_COMPLETER, a.auto)
    ident["telephone"] = demander("Téléphone", ident["telephone"], a.auto)
    ident["ville"] = demander("Ville", ident["ville"], a.auto)
    # Metier : la ligne sous le nom est souvent le titre du CV.
    lignes = [l.strip() for l in txt.splitlines() if l.strip()]
    propose = next((l for l in lignes[1:8] if 4 < len(l) < 60 and not re.search(r"[@\d]", l)), "")
    metiers = [m.strip() for m in demander(
        "Métiers visés, séparés par des virgules (ex. : Chargé de communication, Community manager)",
        propose or A_COMPLETER, a.auto, a.metiers).split(",") if m.strip()]
    try:
        annees = int(demander("Années d'expérience professionnelle (stages compris)", "0", a.auto, a.annees))
    except ValueError:
        annees = 0
    # Un intitule au feminin (« Chargée », « Directrice ») propose « f ».
    fem = bool(re.search(r"\b\w+(?:ée|trice|euse)\b", " ".join(metiers), re.I))
    genre = demander("Accords dans les lettres : « diplômé » (m) ou « diplômée » (f) ?",
                     "f" if fem else "m", a.auto, a.genre)
    import zones
    # Le secteur le plus PRECIS qui nomme la ville : la metropole avant la region.
    v = _sans_accent(ident["ville"] or "").lower()
    candidats = [(len(s.get("departements", ())), cle) for cle, s in zones.CATALOGUE.items()
                 if v and v in s.get("lieux", ())]
    secteur_ville = min(candidats)[1] if candidats else ""
    print(f"     secteurs connus : {', '.join(zones.CATALOGUE)}")
    prio = [x.strip() for x in demander("Secteurs PRIORITAIRES (séparés par des virgules)",
                                        secteur_ville or "idf", a.auto, a.prioritaires).split(",") if x.strip()]
    sec = [x.strip() for x in demander("Secteurs secondaires (vide = aucun)", "", a.auto, a.secondaires).split(",")
           if x.strip()]
    inconnus = [x for x in prio + sec if x not in zones.CATALOGUE]
    if inconnus:
        print(f"  ! secteurs inconnus ignores : {inconnus} (voir zones.CATALOGUE, ou les declarer "
              "a la main dans le profil)")
        prio = [x for x in prio if x in zones.CATALOGUE]
        sec = [x for x in sec if x in zones.CATALOGUE]
    ordre, _ = familles_visees(metiers, comps)
    finance = bool(ordre) and ordre[0] in ("finance", "rse", "conseil")
    garder = demander("Garder les ~90 entreprises de finance / ESG suivies par défaut ? (o/n)",
                      "o" if finance else "n", a.auto, a.entreprises).lower().startswith("o")

    rep = {"identite": ident, "competences": comps, "metiers": metiers, "annees": annees,
           "genre": "f" if genre.lower().startswith("f") else "m",
           "prioritaires": prio, "secondaires": sec}

    print("\n3. Ecriture")
    cv_dir = RACINE / "CV" / "Modèles"
    cv_dir.mkdir(parents=True, exist_ok=True)
    prefixe = slug(f"{ident['prenom']} {ident['nom']}")
    cv_nom = f"CV_{prefixe}{cv.suffix.lower()}"
    shutil.copyfile(cv, cv_dir / cv_nom)
    print(f"  CV modele     : CV/Modèles/{cv_nom}")

    profil = construire(txt, cv_nom, rep)
    sortie.parent.mkdir(parents=True, exist_ok=True)
    if sortie.exists():
        shutil.copyfile(sortie, sortie.with_suffix(".json.bak"))
    sortie.write_text(json.dumps(profil, ensure_ascii=False, indent=1), encoding="utf-8")
    # Les familles du tri se calculent avec le profil en place (tri_file le lit).
    import os
    import profil as _profil
    os.environ["PROFIL_JOBHUNTER"] = str(sortie)
    _profil.FICHIER = sortie
    _profil.charger.cache_clear()
    profil["tri"]["metiers_exclus"] = familles_tri_a_exclure(metiers)
    sortie.write_text(json.dumps(profil, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"  profil        : {sortie.relative_to(RACINE)}")

    # Fichiers de donnees vides : la file d'offres et le suivi des candidatures.
    # L'app repere sa racine a Scripts/apply_queue.json.
    for nom, vide in (("apply_queue.json", "[]"), ("suivi.json", "{}")):
        if not (HERE / nom).exists():
            (HERE / nom).write_text(vide, encoding="utf-8")
    faits = RACINE / "CV" / "Faits.json"
    if not faits.exists():
        faits.write_text(json.dumps({"_lisez_moi": (
            "Faits VALIDES ajoutes au CV quand une annonce les appelle. Voir "
            "profil/exemple_Faits.json pour le format."), "faits": [],
            "allegements": {"fr": [], "en": []}}, ensure_ascii=False, indent=1), encoding="utf-8")
        print("  faits du CV   : CV/Faits.json (vide)")

    cfg_p = HERE / "sources_config.json"
    cfg = json.loads(cfg_p.read_text(encoding="utf-8")) if cfg_p.exists() else {}
    cfg["keywords"] = profil["collecte"]["mots_cles"]
    if not garder:
        cfg["ats_slugs"], cfg["workday_flux"], cfg["autres_flux"] = {}, [], []
    cfg_p.write_text(json.dumps(cfg, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"  recherche     : {len(cfg['keywords'])} mot(s)-cle(s)"
          + ("" if garder else " · listes d'entreprises finance/ESG videes"))

    print("\n4. A relire dans profil/profil.json")
    manque = []
    if A_COMPLETER in json.dumps(profil["identite"], ensure_ascii=False):
        manque.append("identite : champs « A COMPLETER »")
    if not profil["domaines"][0]["parcours_fr"]:
        manque.append("domaines[].parcours_fr : le paragraphe de parcours de tes lettres (3-4 phrases "
                      "factuelles : formation, experiences, ce que tu apportes) — VIDE pour l'instant")
    if len(comps) < 5:
        manque.append(f"competences : seulement {len(comps)} reconnue(s) — ajoute les tiennes "
                      "(libelle + synonymes que les annonces emploient)")
    manque.append("message.diplome_fr / langues_fr et lettre.phrase_profil_fr : une phrase chacun")
    if cv.suffix.lower() == ".pdf":
        manque.append("CV en PDF : il partira tel quel. Fournis un .docx pour des CV adaptes a chaque offre")
    for m in manque:
        print(f"  - {m}")
    print("\nVerifier : python Scripts/profil.py   ·   Lancer l'app : python JobHunter.pyw")
    return 0


if __name__ == "__main__":
    sys.exit(main())
