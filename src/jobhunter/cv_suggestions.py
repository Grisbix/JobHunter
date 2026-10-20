"""Suggestions from the CV, each with the sentence that justifies it — never applied alone.

Two kinds, kept apart on purpose:

- **skills** the CV attests (a skill of the catalogue, programs.skills_catalog, found in
  the CV's text) that the profile does not list yet. Accepted, they go into the profile's
  `competences`, which the letters and CVs draw on: so only what the CV itself says.
- **searches**: the same skills' names as search keywords (`collecte.mots_cles`).
  Searching a term never adds it to the CV or the letters.

Everything is local: the CV is read on this computer, nothing is sent anywhere. The CV's
text is data, never an instruction.
"""

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class Suggestion:
    """One proposal: its kind, its label, the CV's sentence that justifies it."""

    genre: str  # "competence" or "recherche"
    libelle: str
    phrase: str
    entree: dict[str, Any] | None = None  # the profile entry, for a skill


def lire_cv(chemin: Path) -> str:
    """The text of a .docx or .pdf CV (raises ValueError for another format)."""
    suffixe = chemin.suffix.lower()
    if suffixe == ".pdf":
        import fitz

        with fitz.open(str(chemin)) as d:
            return "\n".join(page.get_text() for page in d)
    if suffixe == ".docx":
        from docx import Document

        d = Document(str(chemin))
        lignes = [x.text for x in d.paragraphs]
        lignes += [c.text for t in d.tables for row in t.rows for c in row.cells]
        return "\n".join(lignes)
    raise ValueError(f"format non pris en charge : {chemin.suffix}")


def _phrase(texte: str, synonyme: str) -> str | None:
    """The CV's sentence (or line) where `synonyme` appears, or None."""
    s = synonyme.lower().strip()
    court = len(s) <= 4
    for morceau in re.split(r"(?<=[.!?])\s+|\n", texte):
        m = morceau.lower()
        if (re.search(r"\b" + re.escape(s) + r"\b", m) if court else s in m) and morceau.strip():
            return " ".join(morceau.split())[:200]
    return None


def suggestions(texte: str, profil: dict[str, Any]) -> list[Suggestion]:
    """What the CV attests that the profile does not have yet, with its proof."""
    from jobhunter.programs.skills_catalog import catalogue

    deja = {k.lower() for k in (profil.get("competences") or {})}
    deja_syn = {
        s.lower()
        for v in (profil.get("competences") or {}).values()
        if isinstance(v, dict)
        for s in v.get("synonymes") or []
    }
    mots = {m.lower() for m in ((profil.get("collecte") or {}).get("mots_cles") or [])}
    out: list[Suggestion] = []
    for libelle, v in catalogue().items():
        if libelle.lower() in deja or any(s.lower() in deja_syn for s in v["synonymes"]):
            continue
        phrase = next((p for s in v["synonymes"] if (p := _phrase(texte, s))), None)
        if phrase is None:
            continue  # not in the CV: never proposed
        entree = {k: v[k] for k in ("synonymes", "en", "prose_fr", "prose_en") if k in v}
        if v.get("technique") is False:
            entree["technique"] = False
        out.append(Suggestion("competence", libelle, phrase, entree))
        if v.get("technique") is not False and libelle.lower() not in mots:
            out.append(Suggestion("recherche", libelle.lower(), phrase))
    return out


def appliquer(profil: dict[str, Any], acceptees: list[Suggestion]) -> dict[str, Any]:
    """A copy of `profil` with the accepted suggestions added, nothing removed."""
    import json

    nouveau = json.loads(json.dumps(profil))
    for s in acceptees:
        if s.genre == "competence" and s.entree is not None:
            nouveau.setdefault("competences", {}).setdefault(s.libelle, s.entree)
        elif s.genre == "recherche":
            mots = nouveau.setdefault("collecte", {}).setdefault("mots_cles", [])
            if s.libelle.lower() not in {m.lower() for m in mots}:
                mots.append(s.libelle)
    return nouveau
