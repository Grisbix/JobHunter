"""Edit a generated CV or letter from the window: its text, paragraph by paragraph.

Part of jobhunter.documents. The window shows each paragraph of the .docx (body and
table cells, in document order) and sends back the ones the user changed. A change
only rewrites the characters that differ, inside the runs that held them: the bold of
a heading or the italics of a date stay where they were.

A paragraph whose text changed on disk since it was read is left alone (counted as
refused), as is one whose text the runs do not hold in full (a hyperlink, a field).
"""

from collections.abc import Iterator, Sequence
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def _paragraphes(doc: Any) -> Iterator[Paragraph]:
    """Every paragraph of the body, tables included, in reading order."""
    corps = doc.element.body
    for enfant in corps.iterchildren():
        tag = enfant.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(enfant, doc)
        elif tag == "tbl":
            vues: set[int] = set()
            for ligne in Table(enfant, doc).rows:
                for cellule in ligne.cells:
                    if id(cellule._tc) in vues:  # a merged cell comes back per column
                        continue
                    vues.add(id(cellule._tc))
                    yield from cellule.paragraphs


def lire(docx: Path) -> list[str]:
    """The text of each paragraph, empty ones included (their place is the index)."""
    return [p.text for p in _paragraphes(Document(str(docx)))]


def modifiables(docx: Path) -> list[bool]:
    """Per paragraph: can its text be rewritten (its runs hold all of it)?"""
    return [
        bool(p.runs) and "".join(r.text for r in p.runs) == p.text
        for p in _paragraphes(Document(str(docx)))
    ]


def _remplacer(par: Paragraph, ancien: str, nouveau: str) -> bool:
    """Rewrites `ancien` into `nouveau` in the runs of `par`; False if it cannot."""
    runs = par.runs
    textes = [r.text for r in runs]
    if "".join(textes) != ancien or not runs:
        return False
    debut = 0
    while debut < min(len(ancien), len(nouveau)) and ancien[debut] == nouveau[debut]:
        debut += 1
    fin = 0
    while (
        fin < min(len(ancien), len(nouveau)) - debut
        and ancien[len(ancien) - 1 - fin] == nouveau[len(nouveau) - 1 - fin]
    ):
        fin += 1
    a_fin, insere = len(ancien) - fin, nouveau[debut : len(nouveau) - fin]
    # The characters [debut, a_fin) go; `insere` goes in the run of the first character
    # it replaces, so it keeps that look. A pure insertion takes the look of the
    # character just before it (as typing does); at the very start, the first run's.
    starts = [sum(len(t) for t in textes[:i]) for i in range(len(textes))]
    repere = debut if a_fin > debut else debut - 1
    hote = 0
    if repere >= 0:
        hote = max(i for i, s in enumerate(starts) if s <= repere < s + len(textes[i]))
    for i, (t, s) in enumerate(zip(textes, starts, strict=True)):
        g = min(max(debut - s, 0), len(t))
        d = min(max(a_fin - s, 0), len(t))
        garde = t[:g] + (insere if i == hote else "") + t[d:]
        if garde != t:
            runs[i].text = garde
    return par.text == nouveau


def modifier(docx: Path, avant: Sequence[str], apres: Sequence[str]) -> tuple[int, int]:
    """Writes the paragraphs that changed; (changed, refused). Saves only if one changed."""
    doc = Document(str(docx))
    pars = list(_paragraphes(doc))
    faits = refus = 0
    for i, (a, n) in enumerate(zip(avant, apres, strict=False)):
        if a == n:
            continue
        if i >= len(pars) or pars[i].text != a or not _remplacer(pars[i], a, n):
            refus += 1
            continue
        faits += 1
    if faits:
        doc.save(str(docx))
    return faits, refus
