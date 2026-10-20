"""Optional rewriting by a local model (Ollama).

Part of jobhunter.documents (formerly Scripts/tailor_application.py, phase 3.5).
"""

import importlib
import os
import sys
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from jobhunter import lazy as _paresseux
from jobhunter import log as _journal
from jobhunter import profile as _P
from jobhunter.documents.common import _de, titre_poste
from jobhunter.network import exiger_http

_log = _journal.journal(__name__)


_ici = sys.modules[__name__]
# Values taken from the profile: computed on first read, not at import
# (jobhunter.lazy). In this module, they are read through `_ici.NAME`.
_PARESSEUX: _paresseux.Fabriques = {}


def __getattr__(nom: str) -> Any:
    return _paresseux.lire(_ici, _PARESSEUX, nom)


_PARESSEUX["PROFILE"] = lambda: importlib.import_module("jobhunter.documents.common").PROFILE


# ---- LLM mode: Claude API (if key) OR local Ollama (free) ----
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://localhost:11434")


OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "llama3.1:8b")


def _llm_prompt(offer: dict[str, Any]) -> str:
    return (
        f"Tu es {_ici.PROFILE['nom']}. {_P.get('lettre.presentation_llm', '') or ''} "
        f"Rédige le CORPS (3 paragraphes, français, sobre, factuel, sans clichés ni superlatifs) d'une lettre de motivation "
        f"pour le poste « {offer['role']} » chez {offer['company']}.\n\n"
        f"Missions de l'offre : {offer['missions'][:600]}\n"
        f"Profil recherché : {offer['profil'][:400]}\n\n"
        "Mets en avant les recoupements réels entre son parcours et l'offre. Pas d'en-tête ni de formule "
        "de politesse (déjà gérés ailleurs). Uniquement les 3 paragraphes du corps."
    )


def _ollama_generate(prompt: str) -> str | None:
    """Calls local Ollama (free). Returns the text or None if unavailable."""
    import json as _json
    import urllib.request

    body = _json.dumps(
        {"model": OLLAMA_MODEL, "prompt": prompt, "stream": False, "options": {"temperature": 0.4}}
    ).encode()
    try:
        req = urllib.request.Request(
            exiger_http(f"{OLLAMA_URL}/api/generate"),
            data=body,
            headers={"Content-Type": "application/json"},
        )
        with urllib.request.urlopen(req, timeout=120) as r:  # nosec B310 : scheme checked by exiger_http
            reponse: str = _json.loads(r.read()).get("response", "")
            return reponse.strip()
    except Exception as e:
        _log.info(
            f"  ! Ollama indispo ({type(e).__name__}) — lance `ollama serve` + `ollama pull {OLLAMA_MODEL}`"
        )
        return None


def llm_rewrite(offer: dict[str, Any], cv_docx: Any, letter_docx: Any) -> bool:
    """Rewrites the letter's body with a language model, if there is one.

    Claude if an ANTHROPIC_API_KEY key is set, otherwise Ollama locally.

    Args:
        offer: the offer (role, company, missions…).
        cv_docx: the generated CV (not changed here).
        letter_docx: the generated letter, whose body is replaced.

    Returns:
        True if the letter was rewritten.
    """
    prompt = _llm_prompt(offer)
    body = None
    # 1) Claude API if key
    key = os.environ.get("ANTHROPIC_API_KEY")
    if key:
        try:
            import anthropic

            msg = anthropic.Anthropic(api_key=key).messages.create(
                model="claude-sonnet-5",
                max_tokens=700,
                messages=[{"role": "user", "content": prompt}],
            )
            body = msg.content[0].text.strip()
            _log.info("  ✎ Claude API")
        except Exception as e:
            _log.warning(f"  ! Claude API échec: {type(e).__name__}")
    # 2) sinon Ollama local (gratuit)
    if not body:
        body = _ollama_generate(prompt)
        if body:
            _log.info(f"  ✎ Ollama local ({OLLAMA_MODEL})")
    if not body:
        return False
    # rebuilds the letter with the LLM body
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for line in [
        _ici.PROFILE["nom"],
        _ici.PROFILE["email"],
        _ici.PROFILE["tel"],
        "",
    ]:  # no address/city
        doc.add_paragraph(line)
    # same cleaning + elision as tailor_letter: this LLM path wrote the raw title
    _t = titre_poste(offer.get("role") or "")
    doc.add_paragraph(
        f"Objet : Candidature au poste {_de(_t)}" if _t else "Objet : Candidature spontanée"
    ).runs[0].bold = True
    doc.add_paragraph("Madame, Monsieur,")
    for para in body.split("\n\n"):
        p = doc.add_paragraph(para.strip())
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph("Je vous prie d'agréer, Madame, Monsieur, mes salutations distinguées.")
    doc.add_paragraph(_ici.PROFILE["nom"])
    doc.save(str(letter_docx))
    return True
