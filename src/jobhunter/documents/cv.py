"""CV per offer: copy of the variant, title, targeted skills, tailor-made touch-ups.

Part of jobhunter.documents (formerly Scripts/tailor_application.py, phase 3.5).
"""

import importlib
import json
import shutil
import sys
from pathlib import Path
from typing import Any

from docx import Document

from jobhunter import lazy as _paresseux
from jobhunter import log as _journal
from jobhunter import paths as chemins
from jobhunter.documents.common import CV_ATS, _de, _slug, titre_poste
from jobhunter.documents.rendering import to_pdf
from jobhunter.scoring.lexicon import cv_variant_key

_log = _journal.journal(__name__)


_ici = sys.modules[__name__]
# Values taken from the profile: computed on first read, not at import
# (jobhunter.lazy). In this module, they are read through `_ici.NAME`.
_PARESSEUX: _paresseux.Fabriques = {}


def __getattr__(nom: str) -> Any:
    return _paresseux.lire(_ici, _PARESSEUX, nom)


_PARESSEUX["CV_OUT"] = lambda: importlib.import_module("jobhunter.documents.common").CV_OUT


_PARESSEUX["CV_VARIANT"] = lambda: importlib.import_module("jobhunter.scoring.lexicon").CV_VARIANT


_PARESSEUX["CV_VARIANT_EN"] = lambda: (
    importlib.import_module("jobhunter.scoring.lexicon").CV_VARIANT_EN
)


def cv_sur_mesure(offer: dict[str, Any]) -> dict[str, Any] | None:
    """The CV touch-ups written by hand for this offer, or None.

    `CV/Sur_mesure/<entreprise>__<poste>.json`, same tolerant name as the letter:
        {"remplace":    [["template text", "text for this offer"], ...],
         "ajoute_apres": [["text of a template paragraph", "new paragraph"], ...],
         "supprime":    ["text of a paragraph to remove", ...]}

    `supprime` exists for the page: adding a line specific to the offer often pushes
    the CV to two pages, and the space is taken back from what that particular offer
    does not read.

    The template stays the source: only what the offer asks for is touched up, and
    regeneration takes the touch-ups back instead of overwriting them.
    """
    from jobhunter.paths import CV_SUR_MESURE

    if not CV_SUR_MESURE.exists():
        return None
    vise = _slug(f"{offer.get('company')}__{offer.get('role')}")
    for f in CV_SUR_MESURE.glob("*.json"):
        if _slug(f.stem) == vise:
            retouches: dict[str, Any] = json.loads(f.read_text(encoding="utf-8"))
            return retouches
    return None


def appliquer_cv_sur_mesure(doc: Any, retouches: dict[str, Any]) -> list[str]:
    """Applies the touch-ups to a Document. Returns the list of anchors not found.

    The replacement is done INSIDE the run that carries the text: an experience line
    mixes a bold title and a tabbed date, and rewriting the whole paragraph would
    lose that formatting. An anchor that cannot be found is REPORTED, never ignored:
    the template may have changed, and a touch-up that no longer applies would send
    the generic CV without a word.
    """
    import copy

    manquees = []
    for ancien, nouveau in retouches.get("remplace", []):
        fait = False
        for p in doc.paragraphs:
            for r in p.runs:
                if ancien in r.text:
                    r.text = r.text.replace(ancien, nouveau)
                    fait = True
            if not fait and ancien in p.text:
                txt = p.text.replace(ancien, nouveau)
                for r in p.runs:
                    r.text = ""
                p.runs[0].text = txt
                fait = True
        if not fait:
            manquees.append(ancien)
    for ancre, texte in retouches.get("ajoute_apres", []):
        p = next((p for p in doc.paragraphs if ancre in p.text), None)
        if p is None:
            manquees.append(ancre)
            continue
        el = copy.deepcopy(p._p)
        p._p.addnext(el)
        from docx.text.paragraph import Paragraph

        neuf = Paragraph(el, p._parent)
        for r in neuf.runs:
            r.text = ""
        neuf.runs[0].text = texte
    for ancre in retouches.get("supprime", []):
        p = next((p for p in doc.paragraphs if ancre in p.text), None)
        if p is None:
            manquees.append(ancre)
            continue
        p._p.getparent().remove(p._p)
    return manquees


# ---- CV: injects title + targeted skills line ----
def tailor_cv(
    offer: dict[str, Any],
    nom: str | None = None,
    langue: str = "fr",
    allege: int = 0,
    auto: bool = True,
) -> Path | None:
    """`nom`: imposed stem.

    Without it, several offers with a similar title produce the SAME file and overwrite
    each other — see generate_docs._noms_uniques.

    `langue`: 'en' picks the English template when it exists. 1525 offers of the
    queue are written in English; sending them a French CV disqualifies you before
    reading. The fallback to French is SILENT when the English template is missing
    (not all variants are translated yet) — a French CV is better than no CV.
    """
    chemins.preparer()  # output folders, at write time
    key = cv_variant_key(offer["domaine"])
    src = CV_ATS / (_ici.CV_VARIANT.get(key) or "")
    if not _ici.CV_VARIANT.get(key):
        # Variant without a file: the first CV of the profile that exists.
        src = next(
            (CV_ATS / f for f in _ici.CV_VARIANT.values() if f and (CV_ATS / f).exists()), src
        )
    if (
        langue == "en"
        and _ici.CV_VARIANT_EN.get(key)
        and (CV_ATS / _ici.CV_VARIANT_EN[key]).exists()
    ):
        src = CV_ATS / _ici.CV_VARIANT_EN[key]
    if not src.is_file():
        _log.info(f"  ! CV variante absente: {src.name or key} (profil.json, domaines[].cv)")
        return None
    from jobhunter.programs.rename_docs import new_name

    if src.suffix.lower() == ".pdf":
        # CV provided as PDF only: it goes as is, without adaptation.
        # A .docx template allows inserting the job title into it.
        out: Path = _ici.CV_OUT / (
            (nom or new_name("CV", offer.get("role") or "", offer["company"])) + ".pdf"
        )
        shutil.copyfile(src, out)
        return out
    doc = Document(str(src))
    # The marker is in the PROFILE: "je suis a la recherche d'un poste de
    # [nom du poste]". Pasting the RAW title there gave "d'ACTUAIRE (F/H)" or
    # "d'Adjoint Risque Clients (F/H) - CDD": the gender mention and the contract
    # type have nothing to do in a sentence written in the first person.
    # On a CV the title goes INTO a sentence: it must read well, not only
    # be exact. "…a la recherche d'un poste d'Analyste RSE-ESG - Analyste
    # Credit et ESG." or "…a position as Consultant, Climate Change and
    # Sustainability Services Financial Services." drag along the whole posting
    # title, separating dash included. We keep the first segment when it
    # is enough to name the job; the LETTER's subject, on the other hand, keeps the full
    # title, because it must match the posting exactly.
    role = titre_poste(offer.get("role") or "") or "ce poste"
    # A TITLE THAT IS TOO LONG PUSHES THE CV ONTO TWO PAGES. It goes INTO
    # a sentence of the profile, which fits on three lines and has only one point of
    # margin. Measured on 2026-09-22 on the queue's ESG titles: two of the
    # three longest pushed the CV to two pages, silently.
    #
    # So we do not only cut on "-": postings also stack
    # commas, parentheses and slashes. "Investment
    # Banking, Natural Resources & Energy Transition (Power/Energy) » devient
    # "Investment Banking", which reads well and fits.
    #
    # The 12-character threshold protects heads too short to name a
    # job: "Stage" or "VIE" alone say nothing.
    for sep in (" - ", " – ", ", ", " (", " / ", " | "):
        tete = role.split(sep)[0].strip()
        if 12 <= len(tete) < len(role):
            role = tete
    # Last safety net, if no separator appears: we cut on a word.
    if len(role) > 46:
        coupe = role[:46]
        i = coupe.rfind(" ")
        role = (coupe[:i] if i > 28 else coupe).rstrip(" ,;:.-&")

    # 1) replaces the title [nom du poste], with the elision of _de() (h and y excluded)
    #    The English template carries "[job title]" and needs NO elision:
    #    "as [job title]" is built without a contracted article.
    for p in doc.paragraphs:
        if "[job title]" in p.text:
            txt = p.text.replace("[job title]", role)
        elif "[nom du poste]" in p.text:
            if "de [nom du poste]" in p.text:
                # TYPOGRAPHIC apostrophe: the template writes "d’un poste".
                txt = p.text.replace("de [nom du poste]", _de(role, "’"))
            else:
                txt = p.text.replace("[nom du poste]", role)
        else:
            continue
        for r in p.runs:
            r.text = ""
        (p.runs[0] if p.runs else p.add_run("")).text = txt

    # 1 bis) Touch-ups written by hand for THIS offer (CV/Sur_mesure).
    #       Without them, the automatic facts of CV/Faits.json (cv_auto):
    #       the tailor-made version, reviewed by the user, always wins.
    retouches = cv_sur_mesure(offer)
    if retouches:
        manquees = appliquer_cv_sur_mesure(doc, retouches)
        for m in manquees:
            _log.info(f"  ! CV sur mesure : texte introuvable dans le modele : « {m[:60]} »")
    elif auto:
        from jobhunter.documents import cv_auto

        faits_auto = cv_auto.retouches(
            offer.get("offer_text") or offer.get("role") or "",
            langue if langue == "en" else "fr",
            allege,
        )
        # Missed anchors: NORMAL here. Each CV variant does not have the same
        # lines, and a fact sometimes targets only one of them.
        appliquer_cv_sur_mesure(doc, faits_auto)

    # 2) REMOVED (2026-07-12): the line "Compétences ciblées pour ce poste : …"
    #    gave away an auto-generated CV to recruiters. The skills still serve
    #    the scoring (fit/ATS) and the LETTER, but are no longer injected into the CV.

    from jobhunter.programs.rename_docs import (
        new_name,  # Prenom_Nom_CV_<Poste court>_<Entreprise>
    )

    sortie: Path = _ici.CV_OUT / ((nom or new_name("CV", role, offer["company"])) + ".docx")
    doc.save(str(sortie))
    return sortie


def _pages_pdf(pdf: Path) -> int:
    try:
        import fitz

        with fitz.open(str(pdf)) as d:
            pages: int = d.page_count
            return pages
    except Exception:
        return 1  # unreadable: no point looping


def cv_une_page(
    offer: dict[str, Any], nom: str | None = None, langue: str = "fr"
) -> tuple[Path | None, Path | None]:
    """tailor_cv + to_pdf, guaranteeing ONE page. Returns (docx, pdf).

    The automatic facts add lines; if the PDF overflows, the `allegements` of
    CV/Faits.json are removed one by one. As a last resort, the CV starts again from
    the template without any automatic fact: a one-page generic CV is better than a
    two-page tailored one. A CV/Sur_mesure is never lightened: it was reviewed, and
    so was its page.
    """
    chemins.preparer()  # output folders, at write time
    from jobhunter.documents import cv_auto

    lg = langue if langue == "en" else "fr"
    essais = [(k, True) for k in range(cv_auto.nb_allegements(lg) + 1)] + [(0, False)]
    if cv_sur_mesure(offer):
        essais = [(0, True)]
    cv = pdf = None
    for allege, auto in essais:
        cv = tailor_cv(offer, nom=nom, langue=langue, allege=allege, auto=auto)
        if not cv:
            return None, None
        pdf = to_pdf(cv)
        if not pdf or _pages_pdf(pdf) <= 1:
            break
    return cv, pdf
