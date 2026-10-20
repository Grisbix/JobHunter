"""Full-screen detail of an offer and CV and letter viewer.

Part of jobhunter.app (former JobHunter.pyw, phase 3.7)."""

import importlib
import subprocess
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import scrolledtext, ttk
from typing import Any

from jobhunter import lazy as _paresseux
from jobhunter.app import inventory as _inventaire
from jobhunter.app import theme as _theme
from jobhunter.app._app_type import AppTypee
from jobhunter.app.common import CV_ATS_DIR, PY, SCRIPTS, ouvrir
from jobhunter.app.language import t as _t
from jobhunter.app.theme import (
    ACC,
    ACC_H,
    ACC_TX,
    BG,
    BTN,
    BTN_H,
    CARD,
    CB,
    CE,
    CF,
    CL,
    LINE,
    MUT,
    ROW0,
    TX,
    WORK,
)
from jobhunter.links import chemin_de

_ici = sys.modules[__name__]
# Values taken from the profile: computed on first read, not at import
# (jobhunter.lazy). In this module, they are read through `_ici.NAME`.
_PARESSEUX: _paresseux.Fabriques = {}


def __getattr__(nom: str) -> Any:
    return _paresseux.lire(_ici, _PARESSEUX, nom)


_PARESSEUX["_PREFIXE"] = lambda: importlib.import_module("jobhunter.app.common")._PREFIXE


_PARESSEUX["CV_OFFRE_DIR"] = lambda: importlib.import_module("jobhunter.app.common").CV_OFFRE_DIR


_PARESSEUX["LETTRES_DIR"] = lambda: importlib.import_module("jobhunter.app.common").LETTRES_DIR


class DetailMixin:
    def _detail_full(self: "AppTypee", o: Any = None) -> None:
        """Détails PLEIN ÉCRAN, fenêtre COUPÉE EN DEUX : infos offre à GAUCHE,
        CV / lettre liés à DROITE (aperçu PDF ou docx ÉDITABLE + 💾)."""
        o = o or self._cur
        if not o:
            self._log("Sélectionne d'abord une offre.\n")
            return
        self._dispo_memo()  # "↶ Retour" goes back to the table as it was
        self._hide_all()
        if self._dv is None:
            self._dv = CF(self.main, bg=WORK)
        for c in self._dv.winfo_children():  # rebuilt per offer
            c.destroy()
        w = self._dv
        top = CF(w, bg=BG)
        top.pack(fill="x")
        CB(
            top,
            text="Offres",
            command=self._show_dash,
            bg=BTN,
            fg=MUT,
            relief="flat",
            cursor="hand2",
            font=(_theme.FONT_UI, 13),
        ).pack(side="left", padx=12, pady=(2, 0))
        CL(
            top,
            text=f"{o.get('entreprise', '')} — {o.get('poste', '')[:70]}",
            bg=BG,
            fg="#b45309",
            font=(_theme.FONT_UI, 16),
        ).pack(side="left", padx=28)

        body = tk.PanedWindow(w, orient="horizontal", bg=LINE, sashwidth=5, bd=0)
        body.pack(fill="both", expand=True, padx=24, pady=(4, 4))

        # ── LEFT: offer information + actions ──
        left = CF(body, bg=BG)
        dtxt = scrolledtext.ScrolledText(
            left,
            bg=ROW0,
            fg=TX,
            insertbackground=TX,
            font=(_theme.FONT_UI, 13),
            relief="flat",
            wrap="word",
            padx=14,
            pady=10,
            state="disabled",
        )
        dtxt.tag_configure("k", foreground=ACC, font=(_theme.FONT_UI, 13))
        dtxt.tag_configure("h", foreground="#b45309", font=(_theme.FONT_UI, 16))
        dtxt.pack(fill="both", expand=True)
        lbt = CF(left, bg=BG)
        lbt.pack(fill="x", pady=(4, 0))
        CB(
            lbt,
            text="Envoi confirmé",
            command=lambda: (self._postuler(o), self._fill_detail(dtxt, o)),
            bg=ACC,
            fg=ACC_TX,
            activebackground=ACC_H,
            relief="flat",
            cursor="hand2",
        ).pack(side="left")
        CB(
            lbt,
            text="Ouvrir l'offre",
            command=lambda: self._open_offer(o),
            bg=BTN,
            fg=TX,
            activebackground=BTN_H,
            relief="flat",
            cursor="hand2",
        ).pack(side="left", padx=12)
        bd = CB(
            lbt,
            text="Description",
            bg=BTN,
            fg=TX,
            activebackground=BTN_H,
            relief="flat",
            cursor="hand2",
        )
        bd.configure(command=lambda: self._fetch_desc(o, dtxt, bd))
        bd.pack(side="left", padx=12)
        self._fill_detail(dtxt, o)
        body.add(left, minsize=320)

        # ── RIGHT: linked CV / letter (PDF preview; docx letter = EDITABLE) ──
        right = CF(body, bg=BG)
        rbar = CF(right, bg=BG)
        rbar.pack(fill="x")
        rinfo = CL(right, text="", bg=BG, fg=MUT, font=(_theme.FONT_UI, 12), anchor="w")
        rinfo.pack(fill="x")
        holder = CF(right, bg=BG)
        holder.pack(fill="both", expand=True)
        rcv = tk.Canvas(holder, bg="#f8fafc", highlightthickness=0)
        rsb = tk.Scrollbar(holder, command=rcv.yview)
        rcv.configure(yscrollcommand=rsb.set)
        rtxt = scrolledtext.ScrolledText(
            holder,
            bg="#f8fafc",
            fg=TX,
            insertbackground=TX,
            font=(_theme.FONT_UI, 13),
            relief="flat",
            wrap="word",
            padx=14,
            pady=10,
        )
        st: dict[str, Any] = {"path": None, "imgs": []}

        def surface(kind: Any) -> None:
            rcv.pack_forget()
            rsb.pack_forget()
            rtxt.pack_forget()
            if kind == "pdf":
                rsb.pack(side="right", fill="y")
                rcv.pack(side="left", fill="both", expand=True)
            else:
                rtxt.pack(fill="both", expand=True)

        def load(path: Any) -> None:
            """Charge le doc à droite. Lettre .docx → texte ÉDITABLE ; sinon PDF empilé."""
            # for editing: prefers the letters' twin .docx
            if (
                path.suffix.lower() == ".pdf"
                and _ici.LETTRES_DIR in path.parents
                and path.with_suffix(".docx").exists()
            ):
                path = path.with_suffix(".docx")
            st["path"] = path
            if path.suffix.lower() == ".pdf":
                try:
                    import base64

                    import fitz

                    doc = fitz.open(str(path))
                    surface("pdf")
                    rcv.delete("all")
                    st["imgs"].clear()
                    y = 8
                    for i in range(doc.page_count):
                        pix = doc[i].get_pixmap(matrix=fitz.Matrix(1.25, 1.25))
                        img = tk.PhotoImage(data=base64.b64encode(pix.tobytes("png")))
                        st["imgs"].append(img)
                        rcv.create_image(
                            max(rcv.winfo_width() // 2, pix.width // 2 + 8),
                            y,
                            image=img,
                            anchor="n",
                        )
                        y += pix.height + 12
                    rcv.configure(scrollregion=(0, 0, rcv.winfo_width(), y))
                    rinfo.configure(
                        text=f"{path.name} · {doc.page_count} p. · lecture (PDF) — molette pour dérouler"
                    )
                    doc.close()  # releases the file (images already rendered)
                except Exception as ex:
                    surface("txt")
                    rtxt.delete("1.0", "end")
                    rtxt.insert("end", f"aperçu KO: {type(ex).__name__}")
                return
            surface("txt")
            rtxt.delete("1.0", "end")
            try:
                if path.suffix.lower() == ".docx":
                    from docx import Document

                    rtxt.insert("end", "\n".join(p.text for p in Document(str(path)).paragraphs))
                else:
                    rtxt.insert("end", path.read_text(encoding="utf-8", errors="replace"))
                editable = (path.suffix.lower() == ".txt") or (_ici.LETTRES_DIR in path.parents)
                rinfo.configure(
                    text=f"{path.name} · "
                    + (
                        "✏ ÉDITABLE — modifie puis 💾 (PDF resynchronisé)"
                        if editable
                        else "lecture seule (CV : 📂 Ouvrir → Word)"
                    )
                )
            except Exception as ex:
                rtxt.insert("end", f"lecture KO: {type(ex).__name__}")

        def save_doc() -> None:
            p = st["path"]
            if not p:
                return
            ext = p.suffix.lower()
            if not ((ext == ".txt") or (ext == ".docx" and _ici.LETTRES_DIR in p.parents)):
                self._log("✗ édition réservée aux lettres/.txt — CV : 📂 Ouvrir (Word)\n")
                return
            txt = rtxt.get("1.0", "end").rstrip("\n")
            try:
                if ext == ".txt":
                    p.write_text(txt, encoding="utf-8")
                else:
                    from docx import Document
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    from docx.oxml.ns import qn

                    doc = Document(str(p))
                    lines = txt.split("\n")
                    paras = doc.paragraphs
                    if len(lines) == len(paras):
                        for para, line in zip(paras, lines):
                            if para.text != line:
                                for r in para.runs:
                                    r.text = ""
                                (para.runs[0] if para.runs else para.add_run("")).text = line
                    else:
                        for para in list(paras):
                            para._element.getparent().remove(para._element)
                        for line in lines:
                            np = doc.add_paragraph(line)
                            if len(line.strip()) > 90:
                                np.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    s = doc.settings.element
                    if s.find(qn("w:autoHyphenation")) is None:
                        s.append(s.makeelement(qn("w:autoHyphenation"), {qn("w:val"): "true"}))
                    doc.save(str(p))
                self._log(f"💾 enregistré : {p.name}\n")
            except Exception as ex:
                self._log(f"! sauvegarde KO: {type(ex).__name__}\n")
                return
            tw = p.with_suffix(".pdf")
            if ext == ".docx" and tw.exists():

                def regen() -> None:
                    for exe in (
                        r"C:\Program Files\LibreOffice\program\soffice.exe",
                        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                    ):
                        if Path(exe).exists():
                            try:
                                subprocess.run(
                                    [
                                        exe,
                                        "--headless",
                                        "--convert-to",
                                        "pdf",
                                        "--outdir",
                                        str(p.parent),
                                        str(p),
                                    ],
                                    capture_output=True,
                                    timeout=90,
                                )
                                self.root.after(
                                    0, lambda: self._log(f"📄 PDF resynchronisé : {tw.name}\n")
                                )
                            except Exception:
                                pass
                            return

                threading.Thread(target=regen, daemon=True).start()

        rcv.bind("<MouseWheel>", lambda e: rcv.yview_scroll(-1 if e.delta > 0 else 1, "units"))
        cvp, ltp = self._cv_of(o), self._lettre_of(o)
        if cvp:
            CB(
                rbar,
                text="CV",
                command=lambda: load(cvp),
                bg=BTN,
                fg=TX,
                relief="flat",
                cursor="hand2",
            ).pack(side="left", padx=(0, 2))
        if ltp:
            CB(
                rbar,
                text="Lettre",
                command=lambda: load(ltp),
                bg=BTN,
                fg=TX,
                relief="flat",
                cursor="hand2",
            ).pack(side="left", padx=6)
        CB(
            rbar,
            text="Enregistrer",
            command=save_doc,
            bg=BTN,
            fg=TX,
            activebackground=BTN_H,
            relief="flat",
            cursor="hand2",
        ).pack(side="left", padx=24)
        CB(
            rbar,
            text="Ouvrir",
            command=lambda: st["path"] and ouvrir(st["path"]),
            bg=BTN,
            fg=TX,
            relief="flat",
            cursor="hand2",
        ).pack(side="left", padx=6)
        if not (cvp or ltp):
            rinfo.configure(text=_t("aucun CV/lettre lié à cette offre"))
        body.add(right, minsize=380)
        # opens the LETTER by default (editable), otherwise the CV
        if ltp:
            load(ltp)
        elif cvp:
            load(cvp)
        self._dv.pack(fill="both", expand=True)

    def _cv_viewer(self: "AppTypee", preselect: Any = None) -> None:
        self._hide_all()
        if self.viewer is not None:  # already built → shown again
            self.viewer.pack(fill="both", expand=True)
            if preselect:
                p = Path(preselect)
                if p.exists():
                    self._viewer_show(p)
            return
        self.viewer = w = CF(self.main, bg=WORK)
        w.pack(fill="both", expand=True)
        top = CF(w, bg=BG)
        top.pack(fill="x")
        CB(
            top,
            text="Offres",
            command=self._show_dash,
            bg=BTN,
            fg=MUT,
            relief="flat",
            cursor="hand2",
            font=(_theme.FONT_UI, 13),
        ).pack(side="left", padx=12, pady=(2, 0))
        left = CF(w, bg=CARD, width=340)
        left.pack(side="left", fill="y", padx=(8, 4), pady=8)
        left.pack_propagate(False)
        right = CF(w, bg=BG)
        right.pack(side="left", fill="both", expand=True, padx=(4, 8), pady=8)

        CL(left, text="Chercher une offre (entreprise, poste)", bg=CARD, fg=MUT).pack(
            anchor="w", padx=24, pady=(8, 0)
        )
        srch = CE(left, bg=ROW0, fg=TX, insertbackground=TX, relief="flat")
        srch.pack(fill="x", padx=24, ipady=3)
        lb = tk.Listbox(
            left,
            bg=ROW0,
            fg=TX,
            relief="flat",
            exportselection=False,
            font=(_theme.FONT_UI, 13),
            selectbackground=BTN_H,
        )
        sb = ttk.Scrollbar(left, command=lb.yview)
        lb.configure(yscrollcommand=sb.set)
        lb.pack(side="left", fill="both", expand=True, padx=(8, 0), pady=8)
        sb.pack(side="right", fill="y", pady=24)

        files: list[Any] = []  # aligned with lb: Path, Folder or None (group title)
        MAX_LIGNES = 300

        def fill(flt: Any = "") -> None:
            """Base CV, then one folder per offer (CV and letter together), best ranked first."""
            lb.delete(0, "end")
            files.clear()
            base = sorted(CV_ATS_DIR.glob("*.pdf")) if CV_ATS_DIR.exists() else []
            if base and not flt:
                lb.insert("end", f"── {_t('CV de base')} ({len(base)}) ──")
                files.append(None)
                for f in base:
                    lb.insert(
                        "end", "  " + f.stem.replace(f"CV_{_ici._PREFIXE}_", "").replace("_", " ")
                    )
                    files.append(f)
            ds = _inventaire.dossiers(self.offers, flt)
            prets = sum(1 for d in ds if d.etat == "prêt")
            lb.insert("end", f"── {_t('Dossiers par offre')} ({len(ds)}, {prets} {_t('prêts')}) ──")
            files.append(None)
            marque = {"prêt": "✓", "PDF à refaire": "⚠", "incomplet": "✗"}
            for d in ds[:MAX_LIGNES]:
                suffixe = "" if d.etat == "prêt" else f"  ({_t(d.etat)})"
                lb.insert("end", f"  {marque.get(d.etat, '·')} {d.titre}{suffixe}")
                files.append(d)
            if len(ds) > MAX_LIGNES:
                lb.insert("end", f"  … {len(ds) - MAX_LIGNES} {_t('autres : affine la recherche')}")
                files.append(None)

        fill()
        srch.bind("<KeyRelease>", lambda e: fill(srch.get().strip().lower()))

        # barre titre + infos
        bar = CF(right, bg=BG)
        bar.pack(fill="x")
        title = CL(
            bar,
            text="Sélectionne un CV ou une lettre à gauche",
            bg=BG,
            fg=TX,
            font=(_theme.FONT_UI, 13),
            anchor="w",
        )
        title.pack(side="left", padx=12)
        pginfo = CL(bar, text="", bg=BG, fg=MUT)
        pginfo.pack(side="right", padx=12)

        # two preview surfaces: canvas (PDF, scrollable) / text (docx-txt, scrollable)
        state: dict[str, Any] = {"doc": None, "path": None, "imgs": [], "zoom": 1.4}
        holder = CF(right, bg=BG)
        holder.pack(fill="both", expand=True, pady=(4, 4))
        cv_canvas = tk.Canvas(holder, bg="#f8fafc", highlightthickness=0)
        vsb = tk.Scrollbar(holder, command=cv_canvas.yview)
        cv_canvas.configure(yscrollcommand=vsb.set)
        txtv = scrolledtext.ScrolledText(
            holder,
            bg="#f8fafc",
            fg=TX,
            font=(_theme.FONT_UI, 13),
            relief="flat",
            wrap="word",
            padx=14,
            pady=10,
        )

        def show_surface(kind: Any) -> None:
            cv_canvas.pack_forget()
            vsb.pack_forget()
            txtv.pack_forget()
            if kind == "pdf":
                vsb.pack(side="right", fill="y")
                cv_canvas.pack(side="left", fill="both", expand=True)
            else:
                txtv.pack(fill="both", expand=True)

        def render() -> None:
            """Renders ALL the pages stacked → the wheel scrolls through the whole document."""
            doc = state["doc"]
            if not doc:
                return
            import base64

            import fitz

            try:
                cv_canvas.delete("all")
                state["imgs"].clear()
                y, wmax = 8, 0
                for i in range(doc.page_count):
                    pix = doc[i].get_pixmap(matrix=fitz.Matrix(state["zoom"], state["zoom"]))
                    img = tk.PhotoImage(data=base64.b64encode(pix.tobytes("png")))
                    state["imgs"].append(img)  # refs anti-GC
                    cw = max(cv_canvas.winfo_width(), pix.width + 16)
                    cv_canvas.create_image(cw // 2, y, image=img, anchor="n")
                    y += pix.height + 12
                    wmax = max(wmax, pix.width)
                cv_canvas.configure(scrollregion=(0, 0, wmax + 16, y))
                pginfo.configure(
                    text=f"{doc.page_count} page(s) · zoom {state['zoom']:.1f} · molette = dérouler"
                )
            except Exception as ex:
                cv_canvas.delete("all")
                cv_canvas.create_text(
                    20, 20, text=f"aperçu KO: {type(ex).__name__}", fill="#dc2626", anchor="nw"
                )

        def show(path: Any) -> None:
            # closes the previous PDF doc (otherwise file locked → LibreOffice regeneration impossible)
            if state.get("doc") is not None:
                try:
                    state["doc"].close()
                except Exception:
                    pass
                state["doc"] = None
            state["path"] = path
            title.configure(text=path.name)
            ext = path.suffix.lower()
            if ext == ".pdf":
                try:
                    import fitz

                    state["doc"] = fitz.open(str(path))
                    show_surface("pdf")
                    render()
                except Exception as ex:
                    show_surface("txt")
                    txtv.delete("1.0", "end")
                    txtv.insert("end", f"lecture KO: {type(ex).__name__}")
                return
            # DOCX / TXT → texte complet scrollable
            state["doc"] = None
            show_surface("txt")
            txtv.delete("1.0", "end")
            try:
                if ext == ".docx":
                    from docx import Document

                    txt = "\n".join(p.text for p in Document(str(path)).paragraphs)
                else:
                    txt = path.read_text(encoding="utf-8", errors="replace")
                txtv.insert("end", txt or "(document vide)")
                editable = (ext == ".txt") or (_ici.LETTRES_DIR in path.parents)
                pginfo.configure(
                    text=f"{len(txt)} car. · "
                    + (
                        "✏ ÉDITABLE — modifie puis 💾 Enregistrer"
                        if editable
                        else "lecture seule (CV : 📂 Ouvrir → Word)"
                    )
                )
            except Exception as ex:
                txtv.insert(
                    "end", f"lecture KO: {type(ex).__name__} — clique « Ouvrir » pour Word."
                )

        def on_sel(e: Any = None) -> None:
            c = lb.curselection()  # type: ignore[no-untyped-call]  # curselection : typeshed : non annotee
            if not c or files[c[0]] is None:
                return
            choix = files[c[0]]
            if isinstance(choix, _inventaire.Dossier):
                state["dossier"] = choix
                piece = choix.cv or choix.lettre
                if piece is None:
                    title.configure(text=f"{choix.titre} — aucune pièce sur le disque")
                    return
                show(piece)
            else:
                state["dossier"] = None
                show(choix)

        def basculer() -> None:
            """Switches from the CV to the letter of the same folder, and back."""
            d = state.get("dossier")
            if d is None:
                self._log("Sélectionne d'abord un dossier d'offre dans la liste.\n")
                return
            autre = d.lettre if state.get("path") == d.cv else d.cv
            if autre is None:
                self._log(f"✗ {d.titre} : pièce manquante ({d.etat})\n")
                return
            show(autre)

        lb.bind("<<ListboxSelect>>", on_sel)

        # ── Built-in EDITING: letters (.docx) + .txt. CV .docx = complex layout → Word.
        def save_doc() -> None:
            p = state["path"]
            if not p:
                return
            ext = p.suffix.lower()
            editable = (ext == ".txt") or (ext == ".docx" and _ici.LETTRES_DIR in p.parents)
            if not editable:
                self._log(
                    "✗ édition intégrée réservée aux lettres/.txt — CV .docx : bouton 📂 Ouvrir (Word)\n"
                )
                return
            txt = txtv.get("1.0", "end").rstrip("\n")
            try:
                if ext == ".txt":
                    p.write_text(txt, encoding="utf-8")
                else:
                    from docx import Document
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    from docx.oxml.ns import qn

                    doc = Document(str(p))
                    lines = txt.split("\n")
                    paras = doc.paragraphs
                    if len(lines) == len(paras):
                        # 1 line = 1 paragraph → in place (styles kept)
                        for para, line in zip(paras, lines):
                            if para.text != line:
                                for r in para.runs:
                                    r.text = ""
                                (para.runs[0] if para.runs else para.add_run("")).text = line
                    else:
                        # structure changed → rebuilt (letter = simple layout)
                        for para in list(paras):
                            para._element.getparent().remove(para._element)
                        for line in lines:
                            np = doc.add_paragraph(line)
                            if len(line.strip()) > 90:
                                np.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    s = doc.settings.element  # automatic hyphenation (justified without gaps)
                    if s.find(qn("w:autoHyphenation")) is None:
                        s.append(s.makeelement(qn("w:autoHyphenation"), {qn("w:val"): "true"}))
                    doc.save(str(p))
                self._log(f"💾 enregistré : {p.name}\n")
            except Exception as ex:
                self._log(f"! sauvegarde KO: {type(ex).__name__}\n")
                return
            # twin PDF regenerated in the background (LibreOffice) to stay in sync
            tw = p.with_suffix(".pdf")
            if ext == ".docx" and tw.exists():

                def regen() -> None:
                    for exe in (
                        r"C:\Program Files\LibreOffice\program\soffice.exe",
                        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                    ):
                        if Path(exe).exists():
                            try:
                                subprocess.run(
                                    [
                                        exe,
                                        "--headless",
                                        "--convert-to",
                                        "pdf",
                                        "--outdir",
                                        str(p.parent),
                                        str(p),
                                    ],
                                    capture_output=True,
                                    timeout=90,
                                )
                                self.root.after(
                                    0, lambda: self._log(f"📄 PDF resynchronisé : {tw.name}\n")
                                )
                            except Exception:
                                pass
                            return

                threading.Thread(target=regen, daemon=True).start()

        def zoom(d: Any) -> None:
            state["zoom"] = max(0.6, min(3.0, state["zoom"] + d))
            render()

        def _wheel(e: Any) -> None:
            cv_canvas.yview_scroll(-1 if e.delta > 0 else 1, "units")

        cv_canvas.bind("<MouseWheel>", _wheel)
        cv_canvas.bind("<Enter>", lambda e: cv_canvas.focus_set())

        bt = CF(right, bg=BG)
        bt.pack(fill="x")
        # Second row: the bulk tools. On one row, at 1320 px, they went off the window.
        bt2 = CF(right, bg=BG)
        bt2.pack(fill="x", pady=(4, 0))
        for txt_, cmd in [
            ("↑ Haut", lambda: cv_canvas.yview_moveto(0)),
            ("↓ Bas", lambda: cv_canvas.yview_moveto(1)),
            ("－", lambda: zoom(-0.2)),
            ("＋", lambda: zoom(0.2)),
        ]:
            CB(bt, text=txt_, command=cmd, bg=BTN, fg=TX, relief="flat", cursor="hand2").pack(
                side="left", padx=6
            )
        CB(
            bt,
            text="Enregistrer",
            command=save_doc,
            bg=BTN,
            fg=TX,
            activebackground=BTN_H,
            relief="flat",
            cursor="hand2",
        ).pack(side="left", padx=24)
        CB(
            bt,
            text="Ouvrir",
            command=lambda: state["path"] and ouvrir(state["path"]),
            bg=ACC,
            fg=ACC_TX,
            relief="flat",
            cursor="hand2",
        ).pack(side="left", padx=6)
        # in the title bar: the bottom bar is already full
        CB(
            bar,
            text="Remplacer mon CV de base…",
            command=self._choisir_cv_de_base,
            bg=BTN,
            fg=TX,
            relief="flat",
            cursor="hand2",
        ).pack(side="right", padx=6, before=pginfo)
        CB(
            bar,
            text="CV ⇄ lettre",
            command=basculer,
            bg=ACC,
            fg=ACC_TX,
            relief="flat",
            cursor="hand2",
        ).pack(side="right", padx=6, before=pginfo)
        CB(
            bt2,
            text="Faits du CV",
            command=lambda: self.run_step(
                [PY, "-m", "jobhunter.documents.cv_auto", "--liste"], "Banque de faits"
            ),
            bg=BTN,
            fg=TX,
            relief="flat",
            cursor="hand2",
        ).pack(side="left", padx=(10, 2))

        # bulk tools: UPDATE date/order/recipient of the letters · rename to my name
        def _tool(script: Any, label: Any, btn: Any) -> None:
            btn.configure(state="disabled", text="En cours : " + label)

            def work() -> None:
                try:
                    r = subprocess.run(
                        [PY, "-m", script],
                        cwd=str(SCRIPTS),
                        capture_output=True,
                        text=True,
                        timeout=1800,
                        encoding="utf-8",
                        errors="replace",
                        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                    )
                    out = ((r.stdout or "") + (r.stderr or "")).strip().splitlines()
                    msg = out[-1][:90] if out else "fini"
                except Exception as e:
                    msg = f"! {type(e).__name__}"

                def done() -> None:
                    btn.configure(state="normal", text=label)
                    self._log(f"✅ {label}: {msg}\n")
                    fill(srch.get().strip().lower())
                    self.load_offers()

                self.root.after(0, done)

            threading.Thread(target=work, daemon=True).start()

        b_maj = CB(
            bt2,
            text="MAJ lettres",
            bg=BTN,
            fg=TX,
            activebackground=BTN_H,
            relief="flat",
            cursor="hand2",
        )
        b_maj.configure(
            command=lambda: _tool("jobhunter.programs.refresh_letters", "🗓 MAJ lettres", b_maj)
        )
        b_maj.pack(side="left", padx=6)
        b_ren = CB(bt2, text="Renommer (mon nom)", bg=BTN, fg=TX, relief="flat", cursor="hand2")
        b_ren.configure(
            command=lambda: _tool("jobhunter.programs.rename_docs", "📛 Renommer (mon nom)", b_ren)
        )
        b_ren.pack(side="left", padx=6)
        CB(bt2, text="Retour", command=self._show_dash, bg=BTN, fg=TX, relief="flat").pack(
            side="right", padx=6
        )
        cv_canvas.bind("<Prior>", lambda e: cv_canvas.yview_scroll(-1, "pages"))
        cv_canvas.bind("<Next>", lambda e: cv_canvas.yview_scroll(1, "pages"))
        self._viewer_show = show  # reused when the panel is shown again with preselect

        def _close_doc() -> None:
            if state.get("doc") is not None:
                try:
                    state["doc"].close()
                except Exception:
                    pass
                state["doc"] = None

        self._viewer_close = _close_doc  # called by _show_dash (unlocks the pdf)

        if preselect:
            p = Path(preselect)
            if p.exists():
                show(p)

    def _doc_of(self: "AppTypee", o: Any, key: Any) -> Any:
        """Disk path of the document linked to the offer (cv/letter, decodes file:///…). None if missing."""
        u = o.get(key) or ""
        if not u:
            return None
        p = chemin_de(u)
        return p if p is not None and p.exists() else None

    def _cv_of(self: "AppTypee", o: Any) -> Any:
        return self._doc_of(o, "cv")

    def _lettre_of(self: "AppTypee", o: Any) -> Any:
        return self._doc_of(o, "lettre")
