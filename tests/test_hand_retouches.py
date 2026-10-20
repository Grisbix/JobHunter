"""A document retouched by hand is never overwritten by a regeneration (documents.generated)."""

from docx import Document

from jobhunter.documents import generated as G


def doc(*paragraphes):
    d = Document()
    for p in paragraphes:
        d.add_paragraph(p)
    return d


def texte(path):
    return [p.text for p in Document(str(path)).paragraphs]


def test_premiere_ecriture_puis_remplacement(tmp_path):
    out = tmp_path / "Lettre.docx"
    G.enregistrer(doc("Madame,", "v1"), out)
    G.enregistrer(doc("Madame,", "v2"), out)  # untouched since: replaced
    assert texte(out) == ["Madame,", "v2"]
    assert not (tmp_path / G.A_COMPARER).exists()


def test_retouche_gardee_nouvelle_a_cote(tmp_path):
    out = tmp_path / "Lettre.docx"
    G.enregistrer(doc("Madame,", "v1", "fin"), out)
    doc("Madame,", "ma phrase à moi", "fin").save(str(out))  # the user's retouch
    assert G.enregistrer(doc("Madame,", "v2", "fin"), out) == out
    assert texte(out) == ["Madame,", "ma phrase à moi", "fin"]
    nouveau = tmp_path / G.A_COMPARER / "Lettre.docx"
    assert texte(nouveau) == ["Madame,", "v2", "fin"]
    assert G.ecart(out, nouveau) == 1


def test_document_d_avant_sans_empreinte_remplace(tmp_path):
    out = tmp_path / "Lettre.docx"
    doc("ancienne").save(str(out))  # generated before 3.7: no fingerprint
    G.enregistrer(doc("nouvelle"), out)
    assert texte(out) == ["nouvelle"]


def test_retouche_par_programme_notee(tmp_path):
    out = tmp_path / "Lettre.docx"
    G.enregistrer(doc("le 1er octobre"), out)
    doc("le 20 octobre").save(str(out))  # refresh_letters: our own edit
    G.noter(out)
    G.enregistrer(doc("v2"), out)
    assert texte(out) == ["v2"]


def test_le_dossier_a_comparer_hors_des_listes(tmp_path):
    out = tmp_path / "Lettre.docx"
    G.enregistrer(doc("v1"), out)
    doc("retouche").save(str(out))
    G.enregistrer(doc("v2"), out)
    assert sorted(p.name for p in tmp_path.glob("*.docx")) == ["Lettre.docx"]


def _en_attente(tmp_path):
    out = tmp_path / "Par_offre" / "A" / "Lettre.docx"
    out.parent.mkdir(parents=True)
    G.enregistrer(doc("Madame,", "v1", "fin"), out)
    doc("Madame,", "ma phrase", "fin").save(str(out))
    G.enregistrer(doc("Madame,", "v2", "fin"), out)
    (e,) = G.en_attente([tmp_path / "Par_offre", tmp_path / "absent"])
    return out, e


def test_en_attente_et_comparer(tmp_path):
    out, e = _en_attente(tmp_path)
    assert e.actuel == out and e.nouveau == out.parent / G.A_COMPARER / out.name
    assert G.comparer(e.actuel, e.nouveau) == [
        ("=", "Madame,"),
        ("-", "ma phrase"),
        ("+", "v2"),
        ("=", "fin"),
    ]


def test_prendre_la_nouvelle_garde_l_ancienne_de_cote(tmp_path):
    out, e = _en_attente(tmp_path)
    garde = G.prendre_nouvelle(e)
    assert texte(out) == ["Madame,", "v2", "fin"]
    assert texte(garde) == ["Madame,", "ma phrase", "fin"]  # never deleted
    assert G.en_attente([tmp_path]) == []
    assert not G.retouche(out)  # generated again: the next regeneration may replace it


def test_garder_la_mienne_met_la_nouvelle_de_cote(tmp_path):
    out, e = _en_attente(tmp_path)
    garde = G.garder_mienne(e)
    assert texte(out) == ["Madame,", "ma phrase", "fin"]
    assert texte(garde) == ["Madame,", "v2", "fin"]
    assert G.en_attente([tmp_path]) == []
    assert G.retouche(out)  # still the user's: a later regeneration proposes again
    # a second set-aside of the same name overwrites nothing
    G.enregistrer(doc("Madame,", "v3", "fin"), out)
    (e2,) = G.en_attente([tmp_path])
    garde2 = G.garder_mienne(e2)
    assert garde2 != garde and texte(garde) == ["Madame,", "v2", "fin"]
