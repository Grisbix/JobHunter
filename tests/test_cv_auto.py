"""CV automatique depuis une banque de faits (cv_auto), sans reseau ni LLM.

Seuls les faits VALIDES sont utilises ; declencheurs et faux positifs ; application
a des modeles .docx construits a la volee ; le sur-mesure l'emporte sur
l'automatique. Banques SYNTHETIQUES. Converti depuis Scripts/test_cv_auto.py.
"""

import pytest
from docx import Document

from jobhunter import chemins
from jobhunter import documents as ta
from jobhunter.documents import cv_auto

# ── la banque minimale : un fait valide, un fait NON valide ────────────────────
B = {
    "faits": [
        {
            "id": "vrai",
            "valide": True,
            "declencheurs": ["solaire"],
            "priorite": 5,
            "fr": [{"op": "ajoute_apres", "ancre": "ANCRE", "texte": "ligne vraie"}],
            "en": [],
        },
        {
            "id": "pas_relu",
            "valide": False,
            "declencheurs": ["solaire"],
            "priorite": 9,
            "fr": [{"op": "ajoute_apres", "ancre": "ANCRE", "texte": "ligne inventee"}],
        },
    ],
    "allegements": {"fr": ["L1", "L2"], "en": ["E1"]},
}


def test_fait_valide_retenu_casse_ignoree():
    assert cv_auto.retouches("Projet SOLAIRE en Espagne", "fr", b=B)["faits"] == ["vrai"]


def test_fait_non_valide_jamais_utilise():
    r = cv_auto.retouches("Projet SOLAIRE en Espagne", "fr", b=B)
    assert all("inventee" not in x[1] for x in r["ajoute_apres"])


def test_aucun_declencheur_aucune_retouche():
    assert cv_auto.retouches("comptable", "fr", b=B)["faits"] == []


@pytest.mark.parametrize(
    ("allege", "attendu"),
    [(0, []), (1, ["L1"]), (9, ["L1", "L2"])],
    ids=["rien retire", "premiere ligne", "au-dela sans erreur"],
)
def test_allegements(allege, attendu):
    assert cv_auto.retouches("x", "fr", allege, b=B)["supprime"] == attendu


def test_fait_sans_version_anglaise():
    assert cv_auto.retouches("solaire", "en", b=B)["faits"] == []


# ── une banque realiste : un fait de secteur, de langue, un fait desactive ─────
BANQUE = {
    "faits": [
        {
            "id": "renouvelables",
            "valide": True,
            "priorite": 5,
            "declencheurs": [r"photovolta", r"\bsolaire", r"\b[ée]olien"],
            "fr": [
                {
                    "op": "ajoute_apres",
                    "ancre": "Construction de modèles",
                    "texte": "Études de projets photovoltaïques, dont un parc en Andalousie",
                }
            ],
            "en": [
                {
                    "op": "ajoute_apres",
                    "ancre": "Built valuation models",
                    "texte": "Studies of solar projects, including a park in Andalusia",
                }
            ],
        },
        {
            "id": "espagnol_avance",
            "valide": True,
            "priorite": 3,
            "declencheurs": [r"espagnol", r"spanish", r"español"],
            "fr": [
                {"op": "remplace", "ancre": "Espagnol (débutant)", "texte": "Espagnol (avancé)"}
            ],
            "en": [
                {"op": "remplace", "ancre": "Spanish (beginner)", "texte": "Spanish (advanced)"}
            ],
        },
        {
            "id": "client_unique",
            "valide": False,
            "priorite": 9,
            "declencheurs": [r"photovolta"],
            "fr": [
                {
                    "op": "ajoute_apres",
                    "ancre": "Construction de modèles",
                    "texte": "Mission chez un client précis",
                }
            ],
        },
    ],
    "allegements": {"fr": ["Ligne facultative 1", "Ligne facultative 2"], "en": ["Optional line"]},
}


def ids(texte):
    return [x["id"] for x in cv_auto.faits_retenus(texte, BANQUE)]


def test_photovoltaique_declenche_renouvelables():
    assert "renouvelables" in ids("Investment analyst, centrales photovoltaïques")


def test_fait_desactive_jamais_retenu():
    assert "client_unique" not in ids("photovoltaïque")


def test_energie_ne_declenche_rien():
    assert ids("Data engineer : faire preuve d'énergie") == []


def test_espanol_declenche_espagnol_avance():
    assert "espagnol_avance" in ids("Langues : Español avanzado requis")


# ── application a des modeles .docx construits a la volee ──────────────────────
ANNONCE = "Analyste en investissement photovoltaïque, espagnol B2/C1"


@pytest.fixture
def banque(monkeypatch):
    vrai = cv_auto.retouches
    monkeypatch.setattr(
        cv_auto,
        "retouches",
        lambda t, l="fr", allege=0, b=None: vrai(t, l, allege, BANQUE if b is None else b),
    )


def modele(dossier, lignes, nom):
    d = Document()
    for texte, gras in lignes:
        d.add_paragraph().add_run(texte).bold = gras
    chemin = dossier / nom
    d.save(str(chemin))
    return chemin


@pytest.fixture
def modele_fr(tmp_path):
    return modele(
        tmp_path,
        [
            ("Analyste Financier Junior : Acme Finance", True),
            ("Construction de modèles de valorisation par DCF", False),
            ("Ligne facultative 1", False),
            ("Ligne facultative 2", False),
            ("Espagnol (débutant)", False),
        ],
        "fr.docx",
    )


@pytest.fixture
def cv_fr(banque, modele_fr):
    d = Document(str(modele_fr))
    manquees = ta.appliquer_cv_sur_mesure(d, cv_auto.retouches(ANNONCE, "fr"))
    return d, manquees, [p.text for p in d.paragraphs]


def test_fr_ligne_du_fait_ajoutee(cv_fr):
    assert "parc en Andalousie" in "\n".join(cv_fr[2])


def test_fr_ajoutee_juste_apres_son_ancre(cv_fr):
    textes = cv_fr[2]
    i_fait = textes.index(next(x for x in textes if "Andalousie" in x))
    i_ancre = textes.index(next(x for x in textes if x.startswith("Construction de modèles")))
    assert i_fait == i_ancre + 1


def test_fr_espagnol_avance(cv_fr):
    t = "\n".join(cv_fr[2])
    assert "Espagnol (avancé)" in t and "Espagnol (débutant)" not in t


def test_fr_toutes_les_ancres_trouvees(cv_fr):
    assert cv_fr[1] == []


def test_fr_le_titre_garde_son_gras(cv_fr):
    titre = next(p for p in cv_fr[0].paragraphs if p.text.startswith("Analyste Financier Junior"))
    assert all(r.bold for r in titre.runs if r.text.strip())


@pytest.fixture
def cv_en(banque, tmp_path):
    m = modele(
        tmp_path,
        [
            ("Junior Financial Analyst: Acme Finance", True),
            ("Built valuation models using DCF", False),
            ("Spanish (beginner)", False),
        ],
        "en.docx",
    )
    d = Document(str(m))
    ta.appliquer_cv_sur_mesure(d, cv_auto.retouches(ANNONCE + " Spanish", "en"))
    return "\n".join(p.text for p in d.paragraphs)


def test_en_ligne_du_fait_en_anglais(cv_en):
    assert "park in Andalusia" in cv_en


def test_en_spanish_advanced(cv_en):
    assert "Spanish (advanced)" in cv_en and "Spanish (beginner)" not in cv_en


def test_allegements_deux_lignes_retirees(banque, modele_fr):
    d = Document(str(modele_fr))
    avant = len([p for p in d.paragraphs if p.text.strip()])
    ta.appliquer_cv_sur_mesure(d, cv_auto.retouches("comptable", "fr", allege=2))
    assert avant - len([p for p in d.paragraphs if p.text.strip()]) == 2


# ── le sur-mesure l'emporte sur l'automatique ──────────────────────────────────
@pytest.fixture
def sur_mesure(tmp_path, monkeypatch):
    monkeypatch.setattr(chemins, "CV_SUR_MESURE", tmp_path)
    (tmp_path / "SOLAIRE SA__Analyste Investissement (H-F).json").write_text(
        '{"remplace": [["DCF", "DCF et LBO"]]}', encoding="utf-8"
    )


def test_offre_avec_retouches_a_son_cv_sur_mesure(sur_mesure):
    assert ta.cv_sur_mesure(dict(company="SOLAIRE SA", role="Analyste Investissement (H/F)"))


def test_offre_quelconque_sans_cv_sur_mesure(sur_mesure):
    assert ta.cv_sur_mesure(dict(company="X", role="Y")) is None
