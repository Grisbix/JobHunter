"""Editing a generated CV or letter from the window (jobhunter.documents.editing)."""

from docx import Document

from jobhunter.documents import editing


def _docx(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Objet : ").bold = True
    p.add_run("candidature au poste d'analyste")
    doc.add_paragraph("")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "2024"
    t.cell(0, 1).text = "Analyste ESG, Lyon"
    doc.add_paragraph("Bien cordialement")
    f = tmp_path / "Lettre.docx"
    doc.save(str(f))
    return f


def test_lire_dans_l_ordre_tableaux_compris(tmp_path):
    assert editing.lire(_docx(tmp_path)) == [
        "Objet : candidature au poste d'analyste",
        "",
        "2024",
        "Analyste ESG, Lyon",
        "Bien cordialement",
    ]


def test_modifier_garde_la_mise_en_forme(tmp_path):
    f = _docx(tmp_path)
    avant = editing.lire(f)
    apres = list(avant)
    apres[0] = "Objet : candidature au poste de chargé de mission"
    apres[3] = "Analyste ESG, Paris"
    assert editing.modifier(f, avant, apres) == (2, 0)
    assert editing.lire(f) == apres
    runs = Document(str(f)).paragraphs[0].runs
    assert (runs[0].text, runs[0].bold) == ("Objet : ", True)  # the bold label stays
    assert runs[1].text == "candidature au poste de chargé de mission" and not runs[1].bold


def test_le_texte_tape_prend_l_aspect_d_avant(tmp_path):
    f = _docx(tmp_path)
    avant = editing.lire(f)
    apres = list(avant)
    apres[0] = "Objet : Candidature au poste d'analyste"  # first char of the second run
    assert editing.modifier(f, avant, apres) == (1, 0)
    runs = Document(str(f)).paragraphs[0].runs
    assert [r.text for r in runs] == ["Objet : ", "Candidature au poste d'analyste"]
    apres2 = list(apres)
    apres2[0] = "Nouvel objet"  # everything replaced: the first run's look
    assert editing.modifier(f, apres, apres2) == (1, 0)
    assert editing.lire(f)[0] == "Nouvel objet"


def test_texte_change_sur_le_disque_refuse(tmp_path):
    f = _docx(tmp_path)
    avant = editing.lire(f)
    perime = list(avant)
    perime[4] = "Cordialement"  # not what the file holds
    apres = list(perime)
    apres[4] = "Salutations"
    assert editing.modifier(f, perime, apres) == (0, 1)
    assert editing.lire(f) == avant


def test_un_paragraphe_vide_ne_se_modifie_pas(tmp_path):
    assert editing.modifiables(_docx(tmp_path)) == [True, False, True, True, True]


def test_api_voir_et_modifier_un_document(monkeypatch, tmp_path):
    from jobhunter import root
    from jobhunter.documents import rendering
    from jobhunter.web.api import Api

    (tmp_path / "CV").mkdir()
    f = _docx(tmp_path)
    cv = f.replace(tmp_path / "CV" / "CV_offre.docx")
    (tmp_path / "CV" / "CV_offre.pdf").write_bytes(b"%PDF-1.4")
    monkeypatch.setattr(root, "racine", lambda: tmp_path)
    monkeypatch.setattr(root, "dossier", lambda nom: tmp_path / "Letters")
    monkeypatch.setattr(rendering, "to_pdf", lambda p: p.with_suffix(".pdf"))
    api = Api()
    d = api.document(cv.with_suffix(".pdf").as_uri())
    assert d["ok"] and d["pdf"] == "JVBERi0xLjQ=" and d["paragraphes"][4] == "Bien cordialement"
    apres = list(d["paragraphes"])
    apres[4] = "Bien à vous"
    r = api.enregistrer_document(str(cv), d["paragraphes"], apres)
    assert r == {"ok": True, "modifies": 1, "refuses": 0, "pdf": True}
    assert editing.lire(cv)[4] == "Bien à vous"
    assert editing.lire(cv.with_name("CV_offre.docx.orig"))[4] == "Bien cordialement"
    # outside CV/ and Letters/: refused
    assert not api.document(str(tmp_path / "Lettre.docx"))["ok"]
