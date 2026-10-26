"""The whole text of a CV, wherever its template puts it (jobhunter.cv_text).

Measured on 2026-10-26 with fictitious CVs: a Canva or Word "design" template (text
boxes), a Microsoft resume template (content controls) and a table inside a table gave
0 characters, so the first launch said "almost no text read" and wrote no profile; a
name in the page header was lost.
"""

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from jobhunter.cv_text import texte_fichier
from jobhunter.welcome import arguments_profil

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
LIGNES = [
    "Camille Exemple",
    "Analyste ESG junior",
    "camille.exemple@example.com · 06 00 00 00 00 · Bordeaux",
    "COMPÉTENCES",
    "Analyse ESG · CSRD · SFDR · Empreinte carbone · Excel · Python · SQL",
]


def _paras(lignes):
    return "".join(f"<w:p><w:r><w:t>{l}</w:t></w:r></w:p>" for l in lignes)


def _zone_de_texte(tmp_path):
    # As Word saves a text box: the DrawingML version, and a VML copy as fallback.
    d = Document()
    r = d.add_paragraph().add_run()
    xml = (
        f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:w="{W}">'
        f"<mc:Choice><w:drawing><w:txbxContent>{_paras(LIGNES)}</w:txbxContent></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><w:txbxContent>{_paras(LIGNES)}</w:txbxContent></w:pict></mc:Fallback>"
        "</mc:AlternateContent>"
    )
    r._r.append(etree.fromstring(xml))
    d.save(tmp_path / "cv.docx")
    return tmp_path / "cv.docx"


def test_zone_de_texte_lue_une_fois(tmp_path):
    texte = texte_fichier(_zone_de_texte(tmp_path))
    assert texte.count("Camille Exemple") == 1
    assert "SFDR" in texte


def test_controle_de_contenu(tmp_path):
    d = Document()
    sdt = etree.fromstring(
        f'<w:sdt xmlns:w="{W}"><w:sdtContent>{_paras(LIGNES)}</w:sdtContent></w:sdt>'
    )
    d.element.body.insert(0, sdt)
    d.save(tmp_path / "cv.docx")
    assert "camille.exemple@example.com" in texte_fichier(tmp_path / "cv.docx")


def test_tableau_dans_un_tableau_et_ordre(tmp_path):
    d = Document()
    d.add_paragraph("Camille Exemple")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Profil"
    t.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "CSRD, SFDR"
    d.add_paragraph("Fin")
    lignes = [l for l in texte_fichier_docx(d, tmp_path).splitlines() if l]
    assert lignes == ["Camille Exemple", "Profil", "CSRD, SFDR", "Fin"]


def texte_fichier_docx(d, tmp_path):
    d.save(tmp_path / "cv.docx")
    return texte_fichier(tmp_path / "cv.docx")


def test_en_tete_et_pied_de_page(tmp_path):
    d = Document()
    d.sections[0].header.add_paragraph("Camille Exemple")
    d.sections[0].footer.add_paragraph("camille.exemple@example.com")
    d.add_paragraph("Analyste ESG junior")
    lignes = [l for l in texte_fichier_docx(d, tmp_path).splitlines() if l]
    assert lignes == ["Camille Exemple", "Analyste ESG junior", "camille.exemple@example.com"]


def test_sans_en_tete_rien_n_est_cree(tmp_path):
    d = Document()
    d.add_paragraph("x")
    texte_fichier_docx(d, tmp_path)
    assert Document(str(tmp_path / "cv.docx")).sections[0].header.is_linked_to_previous


def test_autre_format_refuse(tmp_path):
    (tmp_path / "cv.odt").write_bytes(b"")
    with pytest.raises(ValueError):
        texte_fichier(tmp_path / "cv.odt")


def test_reponses_du_formulaire_en_options():
    args = arguments_profil(
        "cv.docx",
        {
            "prenom": " Camille ",
            "nom": "Exemple",
            "email": "",
            "metiers": "Analyste ESG",
            "annees": "2",
            "genre": "f",
            "prioritaires": ["bordeaux", "idf"],
            "secondaires": [],
            "entreprises": True,
        },
    )
    assert args[:2] == ["cv.docx", "--auto"]
    o = dict(zip(args[2::2], args[3::2], strict=True))
    assert o["--prenom"] == "Camille" and o["--email"] == "" and o["--annees"] == "2"
    assert o["--prioritaires"] == "bordeaux,idf" and o["--secondaires"] == ""
    assert (o["--genre"], o["--entreprises"]) == ("f", "o")


def test_profil_d_un_cv_en_zones_de_texte_sans_clavier(tmp_path):
    # The welcome window's path: every answer as an option, no console, no keyboard.
    for d in ("data", "profile"):
        (tmp_path / d).mkdir()
    cv = _zone_de_texte(tmp_path)
    # Enough text to pass the "almost nothing read" check, as a real CV.
    d = Document(str(cv))
    for _ in range(4):
        d.add_paragraph("Stage de six mois : classification SFDR de 300 fonds, reporting CSRD.")
    d.save(str(cv))
    env = {k: v for k, v in os.environ.items() if k != "PROFIL_JOBHUNTER"}
    env.update(JOBHUNTER_RACINE=str(tmp_path), JOBHUNTER_PRINCIPALE=str(tmp_path))
    rep = {
        "prenom": "Camille",
        "nom": "Exemple",
        "email": "camille.exemple@example.com",
        "metiers": "Analyste ESG",
        "annees": 1,
        "genre": "f",
        "prioritaires": ["bordeaux"],
    }
    r = subprocess.run(
        [sys.executable, "-m", "jobhunter.programs.first_launch", *arguments_profil(str(cv), rep)],
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        stdin=subprocess.DEVNULL,
        timeout=120,
        cwd=str(Path(__file__).resolve().parent.parent),
    )
    assert r.returncode == 0, r.stdout + r.stderr
    profil = json.loads((tmp_path / "profile" / "profile.json").read_text(encoding="utf-8"))
    assert profil["identite"]["prenom"] == "Camille"
    assert profil["identite"]["email"] == "camille.exemple@example.com"
