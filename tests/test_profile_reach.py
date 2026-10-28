"""3.53: what a new profile collects — its specialty, its country, and the key it lacks.

Three fictional CVs (a mechanical engineer in Lyon, a waiter in Bordeaux, a lawyer in
Paris) collected on 28/10/2026: the engineer's 375 offers held DevOps posts (the filter
passed « ingénieur » alone) and 126 offers from abroad; the waiter found one offer in
Bordeaux, France Travail's key being absent.
"""

from jobhunter import api_keys
from jobhunter import profile as P
from jobhunter.programs import first_launch as fl


def test_les_deux_parties_dans_n_importe_quel_ordre():
    rx = P.motif(["ingénieur / engineer + mécanique / mechanical"])
    assert P.cherche(rx, "Ingénieur conception mécanique H/F")
    assert P.cherche(rx, "Senior Mechanical Engineer")
    assert not P.cherche(rx, "Ingénieur DevOps")
    assert not P.cherche(rx, "Technicien mécanique")
    # "c++" is a word, not two parts.
    assert P.cherche(P.motif(["c++"]), "Développeur C++")


def test_une_tete_large_demande_sa_specialite():
    mots = fl.filtre_metier(
        ["engineer", "ingénieur", "mécanique", "mechanical"], ["méthodes", "technicien"]
    )
    assert mots[0] == "engineer / ingénieur + mécanique / mechanical"
    assert "ingénieur" not in mots and "technicien" not in mots
    assert "mécanique" in mots and "méthodes" in mots
    rx = P.motif(mots)
    assert P.cherche(rx, "Chargé d'études mécanique")
    assert not P.cherche(rx, "Ingénieur DevOps / Ingénieure DevOps")


def test_une_tete_precise_reste_seule():
    assert fl.filtre_metier(["avocat", "lawyer"], ["juridique"]) == [
        "avocat",
        "lawyer",
        "juridique",
    ]
    # A broad head without a specialty is all there is.
    assert fl.filtre_metier(["consultant"], []) == ["consultant"]


def test_la_cle_conseillee(tmp_path):
    assert api_keys.conseillee(tmp_path, "fr") == "francetravail"
    assert api_keys.conseillee(tmp_path, "de") == "adzuna"
    assert api_keys.conseillee(tmp_path, "jp") is None
    api_keys.enregistrer(tmp_path, "francetravail", {"client_id": "x", "client_secret": "y"})
    assert api_keys.conseillee(tmp_path, "fr") is None


def test_les_pays_vises_a_l_etranger():
    from jobhunter.sources import countries as c

    assert [c.code(x) for x in ("Suisse", "Switzerland", "ch", "Royaume-Uni", "Atlantis")] == [
        "ch",
        "ch",
        "ch",
        "gb",
        "",
    ]
    assert c.etranger_vise(None)[0] and c.etranger_vise(True)[0]  # profiles before 3.53
    assert c.etranger_vise(False) == (False, frozenset(), [])
    assert c.etranger_vise(["Suisse", "Belgique"])[1] == {"ch", "be"}
    assert "de" in c.etranger_vise(["Europe"])[1] and "fr" not in c.etranger_vise(["Europe"])[1]
    assert c.etranger_vise(["Suisse", "monde"])[0]
    assert fl.reponse_monde("") is False and fl.reponse_monde("partout") is True
    assert fl.reponse_monde("Suisse, Atlantis, Europe") == ["Suisse", "Europe"]


def test_la_fusion_garde_les_pays_demandes(monkeypatch):
    from jobhunter.sources import common

    monkeypatch.setattr(common, "MONDE", False)
    monkeypatch.setattr(common, "PAYS_EN_PLUS", frozenset({"ch"}))
    assert common.en_france("Genève, Suisse") and common.en_france("Lyon")
    assert not common.en_france("Bruxelles, Belgique")


def test_le_formulaire_dit_ou_chercher():
    from jobhunter.welcome import arguments_profil

    rep = {"metiers": "x", "prioritaires": ["bordeaux"], "villes": "Pau, bordeaux"}
    o = dict(zip(*[iter(arguments_profil("cv.docx", rep)[2:])] * 2, strict=True))
    assert o["--prioritaires"] == "bordeaux,Pau" and o["--etranger"] == "n"
    assert "--pays-etrangers" not in o
    rep.update(etranger=True, pays_etrangers="Suisse")
    o = dict(zip(*[iter(arguments_profil("cv.docx", rep)[2:])] * 2, strict=True))
    assert (o["--etranger"], o["--pays-etrangers"]) == ("o", "Suisse")


def test_le_profil_garde_les_pays_et_les_villes(tmp_path):
    import json
    import os
    import subprocess
    import sys
    from pathlib import Path

    from docx import Document

    for d in ("data", "profile"):
        (tmp_path / d).mkdir()
    cv = tmp_path / "cv.docx"
    doc = Document()
    for ligne in ("Camille Exemple", "Ingénieure mécanique", "Lyon"):
        doc.add_paragraph(ligne)
    for _ in range(4):
        doc.add_paragraph("Conception de pièces mécaniques, calculs de structure, CAO.")
    doc.save(str(cv))
    env = {k: v for k, v in os.environ.items() if k != "PROFIL_JOBHUNTER"}
    env.update(JOBHUNTER_RACINE=str(tmp_path), JOBHUNTER_PRINCIPALE=str(tmp_path))
    args = [str(cv), "--auto", "--metiers", "Ingénieure mécanique", "--prioritaires"]
    args += ["Mérignac", "--etranger", "o", "--pays-etrangers", "Suisse, Europe"]
    r = subprocess.run(
        [sys.executable, "-m", "jobhunter.programs.first_launch", *args],
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        stdin=subprocess.DEVNULL,
        timeout=120,
        cwd=str(Path(__file__).resolve().parent.parent),
    )
    assert r.returncode == 0, r.stdout + r.stderr
    profil = json.loads((tmp_path / "profile" / "profile.json").read_text(encoding="utf-8"))
    assert profil["monde"] == ["Suisse", "Europe"]
    assert profil["zones"]["prioritaires"] == ["bordeaux"]  # Mérignac's sector
    assert "Suisse" in [s.get("libelle") for s in profil["zones"]["secondaires"]]
