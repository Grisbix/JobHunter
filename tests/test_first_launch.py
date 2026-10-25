"""The profile built from a CV for someone who is not the reference profile.

Measured on 2026-10-24 with a fictitious aerospace engineer's CV ("Ingénieure
structures aéronautiques", Toulouse): the sorting removed her own job from the
queue (the feminine evaded the "industrie" family, so it was excluded), the title
words kept their plural ("structures" did not find "structure"), the LinkedIn URL of
the header made "Réseaux sociaux" a skill, and the three years the CV states were
proposed as 0, which removed every "confirmé" posting.
"""

from jobhunter.programs import first_launch as fl

CV = """Camille MARTIN
Ingénieure structures aéronautiques
camille.martin@exemple.fr · 06 12 34 56 78 · Toulouse · linkedin.com/in/camille-martin
PROFIL
Ingénieure diplômée, spécialisée en calcul de structures composites. Trois ans
d'expérience chez un équipementier de rang 1.
COMPÉTENCES
CATIA V5, Nastran, Patran, éléments finis, Rédaction de dossiers ATA, Python
LANGUES
Français, Anglais (C1)
"""


def test_masculin():
    assert fl.masculin("ingénieure structures") == "ingénieur structures"
    assert fl.masculin("chargée de communication") == "chargé de communication"
    assert fl.masculin("directrice financière") == "directeur financier"
    assert fl.masculin("technicienne") == "technicien"
    assert fl.masculin("conseillère clientèle") == "conseiller clientèle"


def test_mots_significatifs_radicaux():
    # Masculine and singular: the word START finds both forms. "structures" is too
    # wide alone (a "structure d'accueil"): dropped next to a more specific word.
    assert fl.mots_significatifs("Ingénieure structures aéronautiques") == [
        "ingénieur",
        "engineer",
        "ingenieur",
        "aéronautique",
    ]
    # A word ending in "ss" or "us" keeps its end.
    assert fl.mots_significatifs("Process manager") == ["process", "manager"]


def test_le_metier_vise_n_est_pas_exclu_au_feminin():
    exclues = fl.familles_tri_a_exclure(["Ingénieure structures aéronautiques"])
    assert "industrie" not in exclues
    assert "industrie" in fl.familles_tri_a_exclure(["Analyste ESG"])


def test_annees_du_cv():
    assert fl.annees_du_cv(CV) == 3
    assert fl.annees_du_cv("5 ans d'expérience en audit") == 5
    assert fl.annees_du_cv("10+ years of experience") == 10
    assert fl.annees_du_cv("Diplômée en 2024") == 0


def test_genre_feminin_reconnu():
    assert fl._RX_FEMININ.search("Ingénieure structures")
    assert not fl._RX_FEMININ.search("Ingénieur structures")


def test_url_linkedin_n_est_pas_une_competence():
    comps = fl.competences_du_cv(CV)
    assert "Réseaux sociaux" not in comps
    assert "Calcul de structures" in comps  # nastran, éléments finis
    assert "Conception / CAO" in comps  # catia


def test_competences_libres_de_la_section_du_cv():
    connues = fl.competences_du_cv(CV)
    libres = fl.competences_libres(CV, connues)
    # Not in the catalogue: taken as the CV writes it, its own synonym.
    assert libres["Rédaction de dossiers ATA"]["synonymes"] == ["rédaction de dossiers ata"]
    # Covered by a catalogue skill already found: not twice. Languages: another section.
    assert not {"CATIA V5", "Nastran", "Python", "Français"} & set(libres)


def test_construire_profil_d_une_autre_profession():
    comps = fl.competences_du_cv(CV)
    comps.update(fl.competences_libres(CV, comps))
    rep = {
        "identite": fl.identite_du_cv(CV),
        "competences": comps,
        "metiers": ["Ingénieure structures aéronautiques"],
        "annees": fl.annees_du_cv(CV),
        "genre": "f",
        "prioritaires": ["toulouse"],
        "secondaires": [],
    }
    p = fl.construire(CV, "CV_Camille_Martin.docx", rep)
    assert p["collecte"]["mots_cles"] == [
        "ingénieure structures aéronautiques",
        "ingénieur structures aéronautiques",
    ]
    coeur = p["scoring"]["axe_coeur"]
    assert coeur == "ingenierie"
    # The CV's own skill counts in the core axis, as a technical skill.
    assert "Rédaction de dossiers ATA" in p["scoring"]["axes"][coeur]["competences"]
    assert "technique" not in p["competences"]["Rédaction de dossiers ATA"]
    metier = p["filtre_titre"]["metier"]
    assert len(metier) == len(set(metier))
    # Three years: not the junior settings, "confirmé" stays.
    assert "confirmé" not in p["niveau"]["mots_senior"]
    assert p["tri"]["retirer_seniors"] is False


def test_diplome_et_langues_lus_dans_le_cv():
    assert fl.diplome_du_cv(CV) == ""  # no degree line in this short CV
    cv = CV + "FORMATION\n2023 Diplôme d'ingénieur ISAE-SUPAERO, mécanique\n"
    assert fl.diplome_du_cv(cv) == "titulaire d'un diplôme d'ingénieur ISAE-SUPAERO"
    assert fl.diplome_du_cv("Licence d'économie, Lyon 2") == "titulaire d'une licence d'économie"
    assert fl.diplome_du_cv("MBA - HEC Paris") == "titulaire d'un MBA"
    assert fl.langues_du_cv(CV) == ("Je parle anglais (C1).", "I speak English (C1).")
    assert fl.langues_du_cv("Langues : Français, Espagnol, Allemand (B1)") == (
        "Je parle espagnol et allemand (B1).",
        "I speak Spanish and German (B1).",
    )


def test_ton_de_la_lettre_selon_les_annees():
    rep = {
        "identite": fl.identite_du_cv(CV),
        "competences": fl.competences_du_cv(CV),
        "metiers": ["Ingénieure structures"],
        "genre": "f",
        "prioritaires": [],
        "secondaires": [],
    }
    assert fl.construire(CV, "cv.docx", {**rep, "annees": 3})["lettre"]["ton"] == "experimente"
    assert fl.construire(CV, "cv.docx", {**rep, "annees": 1})["lettre"]["ton"] == ""


def test_mots_generiques_ecartes_du_filtre():
    assert fl.mots_significatifs("Infirmière diplômée d'État") == [
        "infirmier",
        "nurse",
        "infirmière",
    ]
    assert fl.mots_significatifs("Contrôleur de gestion") == [
        "contrôleur",
        "controller",
        "controleur",
    ]
    # Nothing more specific: the generic word stays, the filter is never empty.
    assert fl.mots_significatifs("Chargé de projet") == ["chargé", "projet"]


def test_mots_du_titre_aux_deux_genres():
    # "infirmier" does not find "infirmière": both forms, whichever the CV writes.
    assert fl.mots_significatifs("Conseiller") == ["conseiller", "conseillère"]
    assert fl.mots_significatifs("Directrice financière") == [
        "directeur",
        "directrice",
        "financier",
        "financière",
    ]
    # And the other language: an English CV finds the French postings.
    assert fl.mots_significatifs("Developer") == ["developer", "développeur", "developpeur"]


def test_langues_d_un_cv_en_anglais():
    assert fl.langues_du_cv("LANGUAGES\nFrench, English (C1)") == (
        "Je parle anglais (C1).",
        "I speak English (C1).",
    )


def test_benevolat_reconnu():
    from jobhunter.sources.common import contrat_du_titre

    assert contrat_du_titre("Full-Stack Web Developer | Remote Volunteer") == "Bénévolat"
    assert contrat_du_titre("Trésorier bénévole") == "Bénévolat"
    assert contrat_du_titre("VIE - Volontariat International en Entreprise") == "V.I.E"


def test_cv_pdf_et_tableaux_lus_pour_l_ats(monkeypatch, tmp_path):
    # A PDF CV gave "" (ATS 0 % on every offer); a table-laid .docx lost its cells.
    import pymupdf as fitz
    from docx import Document

    from jobhunter.scoring import fit

    pdf = fitz.open()
    pdf.new_page().insert_text((50, 72), "Infirmière en réanimation, hygiène hospitalière")
    pdf.save(tmp_path / "cv.pdf")
    d = Document()
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Nastran, Patran"
    d.save(tmp_path / "cv.docx")
    monkeypatch.setattr(fit, "CV_ATS", tmp_path)
    monkeypatch.setattr(fit, "CV_VARIANT", {"a": "cv.pdf", "b": "cv.docx"}, raising=False)
    assert "réanimation" in fit.cv_text("a")
    assert "nastran" in fit.cv_text("b")


def test_competences_sante_d_un_cv_infirmier():
    cv = (
        "Infirmière en réanimation. Prise en charge de patients lourds, travail en équipe "
        "pluridisciplinaire, hygiène hospitalière, transmissions."
    )
    comps = fl.competences_du_cv(cv)
    for c in (
        "Urgences et réanimation",
        "Prise en charge des patients",
        "Équipe pluridisciplinaire",
        "Hygiène et prévention des infections",
        "Transmissions et dossier de soins",
    ):
        assert c in comps


def test_cv_anglais_resume_en_anglais():
    cv = (
        "Lucas PETIT\nJunior Full-stack Developer\nSUMMARY\nJunior developer with one year of "
        "experience building web applications with React and the Node.js stack, for the web.\n"
        "SKILLS\nReact, Node.js\n"
    )
    rep = {
        "identite": fl.identite_du_cv(cv),
        "competences": fl.competences_du_cv(cv),
        "metiers": ["Junior Full-stack Developer"],
        "annees": 1,
        "genre": "m",
        "prioritaires": [],
        "secondaires": [],
    }
    d = fl.construire(cv, "cv.docx", rep)["domaines"][0]
    assert d["parcours_fr"] == "" and d["parcours_en"]
    # "Software Engineer" is a developer's posting too.
    assert "software" in d["mots_titre"]


def test_titre_dans_une_phrase():
    assert fl.en_phrase("Infirmière diplômée d'État") == "infirmière diplômée d'État"


def test_ats_compte_les_competences_du_metier_absentes_du_cv(monkeypatch):
    # Built from the CV alone, the profile covered 100 % of every posting.
    from jobhunter import profile as _P
    from jobhunter.scoring import fit

    lire = _P.get
    autres = {"Nastran": {"synonymes": ["nastran"]}}
    monkeypatch.setattr(
        _P, "get", lambda k, d=None: autres if k == "competences_ats" else lire(k, d)
    )
    monkeypatch.setattr(fit, "cv_text_cached", lambda k: "")
    r = fit.score_offer(dict(role="Ingénieur calcul", missions="Nastran", profil="", domaine=""))
    assert r["missing"] == {"Nastran"} and r["pct"] == 0
