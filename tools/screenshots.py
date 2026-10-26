"""The README screenshots, on a fictitious profile and fictitious offers.

    uv run python tools/screenshots.py [--lang en] [--out docs/screenshots]

Builds a throwaway data folder (the example profile, "Camille Exemple", and eight
invented companies), prepares the CV and letter of the best offer with the real
engine (`jobhunter documents`, LibreOffice for the PDF), opens the application in the
chosen language and captures the window:

- `offers.png`: the ranked offers, the first one selected (an application with its
  status history and its kept posting; one offer with a knock-out ⚠);
- `letter.png`: an offer opened in full, its letter on the right;
- `web_home.png`, `web_offers_dark.png`: the web interface (3.27), its home screen in
  light mode and its offers in dark mode. The page runs in Playwright's Chromium, its
  calls answered by the real `jobhunter.web.api.Api` on the same fictitious folder:
  nothing of the screen behind is captured.

Nothing of the user's own data is read: the folder is created in a temporary
directory and `JOBHUNTER_RACINE` points to it. Windows only (screen capture by PIL).
"""

import argparse
import ctypes
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
JOUR = time.strftime("%Y-%m-%d")

_DESC = (
    "{poste} at {entreprise}. You will join the sustainability team and work on ESG "
    "analysis of portfolios, CSRD and SFDR reporting, climate risk and carbon footprint "
    "of investments, taxonomy alignment, and engagement with companies. You will build "
    "ESG data dashboards in Excel and Python and present them to investment committees. "
    "Master's degree in finance or sustainability; first experience in responsible "
    "investment appreciated. Fluent English and French."
)

# Offers that match the example profile less and less: a realistic spread of fits.
DESCRIPTIONS = {
    "Climate Risk Analyst": (
        "Climate Risk Analyst. Model physical and transition climate risks for the bank's "
        "credit portfolio, stress tests, scenario analysis, risk reporting to the risk "
        "committee. Python, SQL. Master's degree in finance or engineering."
    ),
    "ESG Data Analyst": (
        "ESG Data Analyst. Collect and clean ESG data from providers, build dashboards in "
        "Power BI, automate data pipelines in Python and SQL, support the sustainability team."
    ),
    "Responsible Investment Associate": (
        "Responsible Investment Associate. Support portfolio managers in asset management: "
        "investment analysis, financial modelling, company meetings, some ESG integration."
    ),
    "Impact Investment Analyst": (
        "Impact Investment Analyst. Due diligence of private equity deals, financial "
        "modelling, impact measurement. 3+ years of experience in investment banking. "
        "You must have the right to work in the UK."
    ),
    "Marketing Assistant": "Marketing assistant for a retail brand: social media and events.",
}

# entreprise, poste, lieu, contrat, salaire
OFFRES = [
    ("Atelier Nova", "ESG Analyst", "Bordeaux", "CDI", "42-48k EUR"),
    ("Studio Lumen", "Sustainable Finance Analyst", "Paris", "CDI", ""),
    ("Fondation Horizon", "Climate Risk Analyst", "Paris", "CDD", ""),
    ("Maison Verte", "CSRD Reporting Officer", "Bordeaux", "CDI", ""),
    ("Groupe Orion", "Responsible Investment Associate", "Lyon", "CDI", "45k EUR"),
    ("Brightside Capital", "Impact Investment Analyst", "London", "CDI", ""),
    ("Kanaal Partners", "ESG Data Analyst", "Amsterdam", "CDI", ""),
    ("Northwind Retail", "Marketing Assistant", "Paris", "CDI", ""),
]


def preparer(racine: Path) -> None:
    """The fictitious data folder, in the 3.0 layout."""
    for d in ("data", "profile", "My_CV", "CV", "Letters"):
        (racine / d).mkdir(parents=True, exist_ok=True)
    profil = json.loads((REPO / "profile" / "example_finance_esg.json").read_text("utf-8"))
    for lg in ("fr", "en"):
        if lg not in profil.setdefault("langues_parlees", []):
            profil["langues_parlees"].append(lg)
    (racine / "profile" / "profile.json").write_text(
        json.dumps(profil, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    lignes = []
    for i, (entreprise, poste, lieu, contrat, salaire) in enumerate(OFFRES):
        desc = DESCRIPTIONS.get(poste) or _DESC.format(poste=poste, entreprise=entreprise)
        lignes.append(
            {
                "entreprise": entreprise,
                "poste": poste,
                "domaine": "esg",
                "statut": f"Scan {JOUR}",
                "url": f"https://jobs.example.com/{i}",
                "contrat": contrat,
                "localisation": lieu,
                "salaire": salaire,
                "source": "greenhouse:example",
                "canal": "auto",
                "site_kind": "entreprise",
                "description": desc,
                "process": JOUR,
                "vu_le": JOUR,
                "ls_check": JOUR,
                "date_publication": "2026-10-08",
            }
        )
    # What 3.3 adds to the card: every search that found the offer, and what its site
    # changed since.
    lignes[0]["vus_sur"] = [
        "greenhouse:example",
        "linkedin-search:esg analyst",
        "welcometothejungle:finance durable",
    ]
    lignes[0]["changements"] = [
        {"date": JOUR, "champ": "salaire", "avant": "40-45k EUR", "apres": "42-48k EUR"}
    ]
    (racine / "data" / "apply_queue.json").write_text(
        json.dumps(lignes, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    # 3.15: an application with its history, and its posting kept (jobhunter postings).
    suivi = {
        "Atelier Nova|ESG Analyst": {
            "s": "entretien",
            "d": JOUR,
            "h": [["2026-10-01", "postulé"], [JOUR, "entretien"]],
        }
    }
    (racine / "data" / "applications.json").write_text(
        json.dumps(suivi, ensure_ascii=False, indent=1), encoding="utf-8"
    )


_SCORER = """
from jobhunter.scoring.fit import fit_detail
from jobhunter.storage import queue as queue_io
with queue_io.transaction() as q:
    for o in q:
        o["fit"], o["exp_requise"], o["fit_fin"] = fit_detail(
            f"{o['poste']} {o['description']}", o["poste"]
        )
"""


def gabarit_cv(racine: Path) -> None:
    """A one-page base CV for the profile's first field (python-docx), fictitious."""
    import docx

    profil = json.loads((racine / "profile" / "profile.json").read_text("utf-8"))
    nom = profil["domaines"][0]["cv"]
    d = docx.Document()
    d.add_heading("Camille Exemple", 0)
    d.add_paragraph("camille.exemple@example.com · +33 6 00 00 00 00 · Bordeaux")
    d.add_heading("Experience", 1)
    d.add_paragraph("Acme Finance — ESG analyst intern (6 months): SFDR classification of funds.")
    d.add_heading("Education", 1)
    d.add_paragraph("Master in Finance, specialisation in sustainable finance.")
    d.add_heading("Skills", 1)
    d.add_paragraph("ESG analysis · CSRD · SFDR · carbon footprint · Excel · Python · SQL")
    (racine / "CV" / "Templates").mkdir(parents=True, exist_ok=True)
    d.save(str(racine / "CV" / "Templates" / nom))


def python(racine: Path, *args: str) -> None:
    env = {**os.environ, "JOBHUNTER_RACINE": str(racine), "PYTHONIOENCODING": "utf-8"}
    env.pop("JOBHUNTER_PRINCIPALE", None)
    subprocess.run([sys.executable, *args], env=env, cwd=racine, check=True)  # nosec B603


def jobhunter(racine: Path, *args: str) -> None:
    env = {**os.environ, "JOBHUNTER_RACINE": str(racine), "PYTHONIOENCODING": "utf-8"}
    env.pop("JOBHUNTER_PRINCIPALE", None)
    subprocess.run(  # nosec B603 : our own program
        [sys.executable, "-m", "jobhunter", *args], env=env, cwd=racine, check=True
    )


def capturer(racine: Path, langue: str, sortie: Path) -> None:
    """Opens the app on `racine` in `langue` and saves the two screenshots."""
    ctypes.windll.shcore.SetProcessDpiAwareness(2)  # sharp, unscaled capture
    os.environ.update(
        JOBHUNTER_RACINE=str(racine), JOBHUNTER_LANGUE=langue, JOBHUNTER_SANS_FOND="1"
    )
    os.environ.pop("JOBHUNTER_PRINCIPALE", None)
    os.environ.pop("JOBHUNTER_ACTION", None)
    import customtkinter as ctk
    from PIL import ImageGrab

    import jobhunter.app as ui

    root = ctk.CTk()
    app = ui.App(root)
    root.state("normal")
    root.geometry("1440x900+0+0")
    root.attributes("-topmost", True)

    def attendre(n: int = 30) -> None:
        for _ in range(n):
            root.update()
            time.sleep(0.02)

    def photo(nom: str) -> None:
        attendre()
        x, y = root.winfo_rootx(), root.winfo_rooty()
        w, h = root.winfo_width(), root.winfo_height()
        ImageGrab.grab(bbox=(x, y, x + w, y + h), all_screens=True).save(sortie / nom)
        print("  ", sortie / nom)

    attendre()
    premier = app.tree.get_children()[0]
    app.tree.selection_set(premier)
    app.tree.focus(premier)
    app._load_edit()
    photo("offers.png")
    app._detail_full(app.rowmap[premier])
    attendre(80)  # the PDF preview renders in the background
    photo("letter.png")
    root.destroy()


# Stands in for pywebview: window.pywebview.api.<name>(...) goes to the Python Api.
_PONT = """
window.pywebview = { api: new Proxy({}, { get: (_, nom) => (...args) => window.__jh(nom, args) }) };
"""


def capturer_web(racine: Path, sortie: Path) -> None:
    """The web interface's two screenshots, rendered headless (see the module's doc)."""
    os.environ.update(JOBHUNTER_RACINE=str(racine))
    os.environ.pop("JOBHUNTER_PRINCIPALE", None)
    from playwright.sync_api import sync_playwright

    from jobhunter.web.__main__ import STATIC
    from jobhunter.web.api import Api

    api = Api()

    def pont(nom: str, args: list) -> object:
        # JSON round trip: what pywebview would hand the page.
        return json.loads(json.dumps(getattr(api, nom)(*args), default=str))

    with sync_playwright() as p:
        nav = p.chromium.launch()
        page = nav.new_page(viewport={"width": 1440, "height": 900})
        page.expose_function("__jh", pont)
        page.add_init_script(_PONT)
        page.goto((STATIC / "index.html").as_uri())
        page.wait_for_timeout(1500)
        # No see-through window here: the page paints its own background.
        page.evaluate("document.body.classList.remove('natif')")

        def photo(nom: str, mode: str, ecran: str) -> None:
            page.evaluate(
                f"(async () => {{ modeAffiche = '{mode}'; appliquerMode();"
                f" await aller('{ecran}'); }})()"
            )
            page.wait_for_timeout(1200)
            page.evaluate("window.scrollTo(0, 0)")  # a focused field may scroll it
            page.screenshot(path=str(sortie / nom))
            print("  ", sortie / nom)

        photo("web_home.png", "clair", "accueil")
        photo("web_offers_dark.png", "sombre", "offres")
        nav.close()


def main() -> int:
    a = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    a.add_argument("--lang", default="en", help="interface language: en, fr, de, es")
    a.add_argument("--out", default=str(REPO / "docs" / "screenshots"))
    a.add_argument("--garder", action="store_true", help="keep the fictitious folder")
    args = a.parse_args()
    sortie = Path(args.out).resolve()
    sortie.mkdir(parents=True, exist_ok=True)
    racine = Path(tempfile.mkdtemp(prefix="jobhunter_captures_"))
    try:
        preparer(racine)
        gabarit_cv(racine)
        python(racine, "-c", _SCORER)
        jobhunter(racine, "documents", "--top", "1")
        jobhunter(racine, "reposts")  # the Signal column: knock-outs
        jobhunter(racine, "postings")  # the posting kept for the application
        capturer_web(racine, sortie)
        capturer(racine, args.lang, sortie)
    finally:
        if args.garder:
            print("dossier fictif :", racine)
        else:
            shutil.rmtree(racine, ignore_errors=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
